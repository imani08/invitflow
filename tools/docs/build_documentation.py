#!/usr/bin/env python3
"""Build version-controlled InvitaFlow documentation, diagrams and Word exports.

Requires python-docx and Pillow in the dedicated documentation environment.
No application/runtime dependency is changed.
"""
from __future__ import annotations

import re
import html
import json
import textwrap
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
TODAY = date(2026, 10, 7).isoformat()
META = f"- Version: 1.0\n- Status: Living document\n- Last updated: {TODAY}\n- Owner: FOCUS HD ENTREPRISES\n- Product: InvitaFlow"

def doc(title: str, body: str, status: str = "IMPLEMENTED") -> str:
    return f"# {title}\n\n{META}\nImplementation status: **{status}**\n\n{body.strip()}\n"

DOCUMENTS: dict[str, str] = {}

DOCUMENTS["00-product/PRODUCT_VISION.md"] = doc("Vision produit InvitaFlow", """
InvitaFlow est une plateforme web de préparation et de gestion d’invitations événementielles. Le dépôt implémente aujourd’hui des espaces événement, cérémonies, invités, placement, conception graphique, génération d’invitations, RSVP, QR/check-in, paiements, crédits et fonctions d’agence/partenaires à des degrés différents.

## Mission et proposition de valeur

Réunir dans un parcours web les tâches dispersées entre organisateurs : structurer un événement, gérer les participants, préparer un visuel cohérent, distribuer des invitations et suivre les réponses. Le produit vise des usages africains et internationaux, avec Kinshasa comme fuseau par défaut de certains parcours. La conception dite « African-first » est une orientation produit et de marque ; elle ne signifie pas que tous les contenus ou marchés sont déjà localisés.

## Utilisateurs

- **Organisateur individuel** : crée ses événements, invités, designs et lots d’invitations à partir de ses crédits.
- **Agence** : un espace agence, clients, membres, quotas et abonnement existent dans les modèles/contrôleurs Events ; les droits et le parcours complet restent à valider.
- **Partenaire** : attribution, commissions et modèles de payout existent dans Payments ; l’encaissement réel des commissions est incomplet.
- **Administrateur** : vues support/modération, finance et pricing sont intégrées à l’application Web et protégées par rôles. L’application `apps/admin` est un shell de navigation, pas un second backend admin.
- **Invité** : répond publiquement par lien signé à l’invitation ; certains parcours peuvent utiliser un QR.

## Modèle économique visible dans le code

Billing publie des barèmes et packs de crédits ; Payments crée des ordres et tente un checkout via un fournisseur ; Wallet conserve un solde de crédits et un ledger. Les crédits sont des unités d’usage, **pas de l’argent**. Le wallet ne détient pas de valeur monétaire.

## État du produit

Plusieurs tranches fonctionnelles sont présentes. Le README et `docs/REMAINING_WORK.md` distinguent explicitement le code disponible de sa préparation production. Runtime intégré, EasyPay/FlexPay réel, certains droits d’agence/partenaire, procédures de reprise et conformité doivent encore être vérifiés ou complétés.

## Sources du dépôt

`README.md`, `docs/REMAINING_WORK.md`, `apps/web/src/app/`, `services/*/src/`, `services/*/prisma/schema.prisma`.
""")

DOCUMENTS["00-product/PRODUCT_SCOPE.md"] = doc("Périmètre produit et maturité", """
## Implémenté dans le dépôt

OIDC BFF Web, profils, événements et cérémonies, invités/imports, plan de tables, documents de design versionnés, assistance éditoriale, jobs de proposition design, média quarantaine/scan, lots d’invitations, rendu PDF/ZIP, liens signés/QR, RSVP/check-in, notifications, audit, billing, ledger Wallet, abstraction de paiement, espaces agence et données partenaires.

## Partiel

Fonctionnalités d’agence, abonnements/quota, commissions et paiements partenaires, fournisseur externe de paiement, opérations de stockage, observabilité et certaines zones d’administration. Le dépôt contient du code et des tests sur ces domaines, mais cela ne prouve pas leur disponibilité complète en production.

## Non vérifié

Parcours navigateur connecté complet, migrations appliquées dans des bases vivantes, rendu Chromium/Linux, délivrabilité SMTP, paiements sandbox et procédures de restauration réellement exécutées.

## Hors périmètre démontré

Aucune promesse de marketplace complète, de payout automatique, de facture PDF, de disponibilité/SLA, de capacité de charge ou de certification de sécurité n’est faite ici sans preuve dans le dépôt.

## Références

Voir `docs/REMAINING_WORK.md`, `docs/E2E_VALIDATION.md`, les README des services et les schémas Prisma.
""", "PARTIAL")

DOCUMENTS["00-product/FUNCTIONAL_REQUIREMENTS.md"] = doc("Exigences fonctionnelles", """
| ID | Exigence traçable au code | Statut | Source principale |
|---|---|---|---|
| FR-AUTH-001 | Authentifier via OIDC Authorization Code avec PKCE S256, state et nonce | IMPLEMENTED, runtime à vérifier | `apps/web/src/lib/auth-session.ts` |
| FR-AUTH-002 | Refuser la session Web sans email vérifié | IMPLEMENTED | `apps/web/src/lib/auth-session.ts` |
| FR-LEGAL-001 | Exiger l’acceptation légale serveur et en conserver les versions/horodatages | IMPLEMENTED dans le provider | `infrastructure/keycloak/providers/invitaflow-legal-acceptance/` |
| FR-EVENT-001 | Créer, lire, modifier, publier/annuler un événement et gérer ses cérémonies | IMPLEMENTED | `services/events/src/events.controller.ts` |
| FR-GUEST-001 | Gérer invités, groupes, accès cérémonie et imports CSV/XLSX | IMPLEMENTED, runtime à vérifier | `services/guests/src/guests.controller.ts` |
| FR-SEAT-001 | Gérer tables, zones, placements et import de plan | IMPLEMENTED | `services/seating/src/seating.controller.ts` |
| FR-DESIGN-001 | Enregistrer des designs et versions JSON associés à un événement | IMPLEMENTED | `services/designs/src/designs.controller.ts` |
| FR-AI-001 | Créer/suivre des jobs d’assistance et de composition selon le fournisseur configuré | PARTIAL | `services/ai-design/README.md` |
| FR-INV-001 | Créer des lots avec instantanés immuables, réservation de crédits et rendu PDF/ZIP | IMPLEMENTED, runtime à vérifier | `services/invitations/README.md` |
| FR-RSVP-001 | Afficher une invitation publique et accepter une réponse par cérémonie | IMPLEMENTED | `services/invitations/src/invitations.controller.ts` |
| FR-CHECKIN-001 | Contrôler l’accès d’un invité par cérémonie avec unicité d’entrée | IMPLEMENTED | `services/invitations/prisma/schema.prisma` |
| FR-PAY-001 | Créer un ordre à partir du catalogue Billing avec clé d’idempotence | IMPLEMENTED | `services/payments/README.md` |
| FR-PAY-002 | Accorder des crédits uniquement après événement de paiement vérifié | IMPLEMENTED dans les consommateurs, intégration à vérifier | `services/wallet/src/payment-consumer.ts` |
| FR-AGENCY-001 | Administrer workspace, membres, clients et quota agence | PARTIAL | `services/events/prisma/schema.prisma` |
| FR-PARTNER-001 | Tracer attribution et commissions partenaires | PARTIAL | `services/payments/prisma/schema.prisma` |
| FR-MEDIA-001 | Recevoir et vérifier un média avant publication dans le bucket privé prêt | IMPLEMENTED | `services/media/README.md` |
| FR-PROFILE-001 | Lire/modifier le profil et gérer une demande de suppression | IMPLEMENTED | `services/profile/src/profile.controller.ts` |
""")

DOCUMENTS["00-product/NON_FUNCTIONAL_REQUIREMENTS.md"] = doc("Exigences non fonctionnelles", """
| ID | Exigence / contrôle constaté | Statut / limite |
|---|---|---|
| NFR-SEC-001 | OIDC code flow, PKCE S256, state/nonce et session BFF opaque chiffrée côté serveur | Code présent, parcours runtime non vérifié |
| NFR-SEC-002 | Services valident signature/issuer/audience/verified email et owner subject | Présent selon guards; couverture inter-utilisateur E2E requise |
| NFR-SEC-003 | Secrets en variables d’environnement/Compose; aucune valeur `.env` versionnée | `.env.example` contient placeholders; audit runtime/deployment requis |
| NFR-SEC-004 | Upload image borné, quarantine, vérification structure, ClamAV, transcodage Sharp | Code/tests unitaires présents; runtime requis |
| NFR-DATA-001 | Bases et identifiants PostgreSQL séparés par service | Compose/provisioning présent |
| NFR-DATA-002 | Migrations Prisma versionnées par service | Présentes; état dans DB inconnue |
| NFR-RES-001 | Outbox, retries, DLQ et reprises sont employés selon consommateurs | Implémentation non uniforme; vérifier catalogue événements |
| NFR-OBS-001 | Health endpoints, métriques Gateway, Prometheus/Grafana et exporters optionnels | Configuration présente; runtime non vérifié |
| NFR-PERF-001 | Limites de taille/dimension média, quotas de jobs IA et claims bornés | Valeurs code/config; aucun rapport de performance E2E actuel |
| NFR-AVAIL-001 | Aucun SLA/RTO/RPO approuvé détecté | À définir; propositions de reprise explicitement étiquetées |
| NFR-A11Y-001 | Styles et composants comprennent focus/thème; aucun audit exhaustif certifié | NOT VERIFIED |
| NFR-PRIV-001 | Données invités sensibles et journaux filtrés selon allowlist dans Audit | Contrôles code présents; durée de rétention juridique à approuver |

Les objectifs numériques de latence, disponibilité, capacité et rétention ne sont pas inventés : ils sont des éléments ouverts tant qu’une politique approuvée ne figure pas au dépôt.
""", "PARTIAL")

DOCUMENTS["00-product/BUSINESS_RULES.md"] = doc("Règles métier présentes", """
- Chaque service métier possède son propre schéma et login DB. Les IDs propriétaires sont généralement des `ownerSubject` issus de l’identité vérifiée.
- La date d’un événement/cérémonie est stockée comme instant avec fuseau déclaré; Events valide l’offset ISO et sa cohérence avec la zone IANA. Zone par défaut à la création : `Africa/Kinshasa`.
- Les importations invités suivent une étape de parsing/mapping/validation avant commit; les erreurs et limites sont exposées comme données de job.
- Une invitation est produite depuis un instantané figé du design, événement, invité et placement; le rendu asynchrone ne lit pas des données modifiables en cours de génération.
- La génération réserve des crédits par invité, puis consomme les rendus réussis et relâche le reste selon résultat/annulation.
- Les crédits représentent une capacité d’utilisation. Un ledger append-only est la source comptable; les corrections passent par des entrées de reversal, jamais par suppression/édition.
- Un événement de paiement vérifié est une condition à l’octroi du pack; le client ne fournit pas le montant autoritaire.
- L’unicité RSVP est par invitation et cérémonie; l’entrée de check-in est protégée par l’unicité invitation/cérémonie.
- Une URL invitation/QR est signée HMAC et révocable; elle ne doit pas encoder contact ou notes d’invité.
- Acceptation CGU/confidentialité à l’inscription est un contrôle serveur et conserve versions et date UTC.

Sources : services Events, Guests, Invitations, Wallet, Payments et provider d’acceptation Keycloak.
""")

DOCUMENTS["01-architecture/ARCHITECTURE.md"] = doc("Architecture InvitaFlow", """
InvitaFlow est un monorepo pnpm piloté par Turborepo. Le navigateur utilise Next.js `apps/web`, qui assure également les routes BFF. Les requêtes API passent par Gateway NestJS/Fastify vers des services de domaine NestJS. Ceux-ci possèdent des schémas et bases PostgreSQL séparés. RabbitMQ distribue les tâches et événements, Redis porte notamment les sessions serveur, et MinIO stocke les médias et documents générés.

```mermaid
flowchart LR
  browser[Browser] --> web[Next.js Web and BFF]
  web --> kc[Keycloak OIDC]
  web --> gateway[Gateway]
  gateway --> services[Domain services]
  services --> db[(Service-owned PostgreSQL databases)]
  services --> mq[RabbitMQ]
  services --> storage[(Private MinIO buckets)]
  web --> redis[(Redis encrypted sessions)]
  worker[Rendering and AI workers] --> mq
  worker --> storage
```

## Services actifs repérés

Profile, Events, Guests, Seating, Designs, Media, AI Design, Billing, Payments, Wallet, Notifications, Invitations/Rendering, Audit et Analytics. `services/access` existe aussi mais son rôle utilisateur/opérationnel doit être considéré selon ses routes/config effectives avant exposition. `apps/admin` fournit le shell admin; consoles métier sont dans Web.

## Flux et frontières

Le navigateur reçoit un cookie opaque HttpOnly et n’est pas le propriétaire d’un bearer token de session BFF. Gateway relaie l’accès vers les APIs. Les services valident l’identité et la propriété; les opérations interservices utilisent des jetons dédiés ou tokens internes selon les routes. Outbox en base métier vers RabbitMQ évite l’écart transaction/message lorsque cette voie est implémentée.

## Infrastructure locale

`compose.yaml` décrit PostgreSQL, Redis, RabbitMQ, Keycloak, Mailpit, MinIO, Gateway, applications et workers. Traefik, ClamAV et le profil d’observabilité sont des composants configurés. Compose est une stack de développement local et ne vaut pas configuration de production.

## Sources

`README.md`, `compose.yaml`, `apps/gateway/src/main.ts`, contrôleurs, modules, `infrastructure/` et `services/*/prisma/schema.prisma`.
""")

DOCUMENTS["01-architecture/C4_CONTEXT.md"] = doc("C4 contexte système", """
InvitaFlow fournit un workspace d’organisation et d’invitation. L’organisateur prépare événement, invités et design; les invités reçoivent un lien pour consulter et répondre. L’équipe support/finance utilise des vues protégées. Keycloak gère l’identité, le fournisseur de paiement traite le checkout, et MinIO/SMTP/RabbitMQ sont des systèmes techniques externes ou infrastructurels.

Voir `diagrams/sources/system-context.mmd` et ses exports. Les intégrations externes réellement qualifiées sont limitées : Mock local est implémenté; FlexPay est un adaptateur incomplet; l’API EasyPay réelle n’est pas démontrée.
""")

DOCUMENTS["01-architecture/C4_CONTAINERS.md"] = doc("C4 conteneurs et frontières", """
| Conteneur | Responsabilité | Données / dépendances |
|---|---|---|
| Web Next.js | UI, BFF OIDC, appels proxifiés | Redis pour sessions chiffrées, Keycloak, Gateway |
| Admin Next.js | Shell de navigation vers vues admin existantes | Web / API protégées |
| Gateway | Proxy API, request IDs/correlation, headers et métriques | services, Identity token |
| Services métier | Règles de domaine isolées | PostgreSQL dédié par service |
| RabbitMQ | Jobs et événements versionnés | queues durables, retries/DLX selon configuration |
| Workers | Rendu invitations, traitements async, AI selon mode | RabbitMQ, snapshots DB, MinIO |
| Keycloak | OIDC, registration, required actions, rôles | PostgreSQL Keycloak |
| MinIO | Assets/quarantine et rendus privés | comptes IAM par usage |
| Mailpit | SMTP local de développement | port SMTP 1025, UI locale 8025 |
| Prometheus/Grafana/Tempo/Loki | Observabilité optionnelle | profil `observability`, maturité non uniforme |
""")

DOCUMENTS["01-architecture/COMPONENT_ARCHITECTURE.md"] = doc("Architecture des composants", """
Le code partagé est volontairement réduit : `packages/contracts`, `packages/ui`, `packages/observability`, `packages/design-document`, `packages/legal-contract`. Les règles événement/invité/seating/design/paiement ne résident pas dans une bibliothèque commune, mais dans le service propriétaire.

Gateway définit des contrôleurs proxy explicites, vérifie le format Bearer et transmet un contexte de corrélation. Les services revalident le token; les routes internes ont des guards séparés. Web BFF échange le code OIDC et conserve tokens chiffrés dans Redis.

Diagramme : `diagrams/sources/component-architecture.mmd`. Sources : `apps/gateway/src/main.ts`, `apps/web/src/lib/auth-session.ts`, `services/*/src/*module.ts`.
""")

DOCUMENTS["01-architecture/DEPLOYMENT_ARCHITECTURE.md"] = doc("Architecture de déploiement", """
`compose.yaml` décrit la topologie de développement. Les réseaux Docker isolent données, identité, application, stockage, messaging, scan média et observabilité. PostgreSQL/Redis/RabbitMQ ne sont pas publiés comme interfaces publiques dans le guide de dépôt; les interfaces dev sont généralement liées à localhost. Les valeurs, profiles, ports et healthchecks exacts sont la source de vérité.

Image Node des projets, image Keycloak custom (provider légal), provisionneurs one-shot, PostgreSQL init, Keycloak init audiences/legal/email, MinIO bootstrap, renderer et services applicatifs démarrent dans un ordre explicite `depends_on`. Un `depends_on` réussi ne prouve pas qu’une intégration fonctionnelle marche.

Ce Compose n’est pas un manifeste de production : aucune topologie HA/cluster, gestionnaire de secrets externe, stratégie rolling deploy ou SLA n’est documentée comme implémentée.
""")

DOCUMENTS["01-architecture/AUTHENTICATION_ARCHITECTURE.md"] = doc("Architecture d’authentification", """
1. Login construit Authorization Code OIDC avec `openid profile email`, state, nonce et challenge PKCE S256.
2. Keycloak authentifie l’utilisateur et applique les required actions (dont VERIFY_EMAIL selon le realm).
3. Le callback consomme le state, échange le code avec code_verifier, vérifie signature/issuer/audience/nonce/azp/subject et exige un email présent avec `email_verified === true`.
4. La session est chiffrée et conservée dans Redis; le navigateur reçoit seulement un cookie opaque HttpOnly/SameSite.
5. Les routes BFF utilisent l’access token serveur côté Gateway/services; refresh est coordonné.

PKCE, session et destination safe sont dans `apps/web/src/lib/auth-session.ts`; callback dans `apps/web/src/app/api/auth/callback/route.ts`; client/claims dans realm-export. Une exécution Keycloak/navigateur reste nécessaire pour démontrer un token réel.
""")

DOCUMENTS["01-architecture/PAYMENT_ARCHITECTURE.md"] = doc("Architecture de paiement", """
Billing fournit prix/packs/règles versionnés. Payments crée un ordre à instantané, choisit un `PaymentProvider`, tente le checkout et enregistre changements/outbox. Un signal de retour navigateur ne constitue pas la preuve de paiement : la vérification serveur du fournisseur est requise avant `payment.succeeded`. Wallet consomme l’événement avec idempotence pour créditer le ledger.

`MockPaymentProvider` permet le développement. Le contrat marchand FlexPay RDC est indiqué indisponible; l’adapter refuse volontairement les opérations réelles. EasyPay n’est pas établi comme fournisseur effectif dans les README actuels. Voir `services/payments/README.md`, `docs/PAYMENTS_EASYPAY.md`, `services/payments/src/providers/`.
""", "PARTIAL")

DOCUMENTS["01-architecture/INVITATION_RENDERING_ARCHITECTURE.md"] = doc("Architecture de rendu invitation", """
Invitations capture un snapshot immuable de chaque invité, du design révisionné, de l’événement, des cérémonies et des données de placement. La demande passe par l’outbox vers une queue RabbitMQ durable. Le worker Rendering revendique les items et utilise le snapshot (pas les objets éditables) pour rendre un PDF A5, stocké privé dans MinIO. ZIP regroupe les résultats. Les états et retries restent enregistrés en DB pour reprise après worker.

Le lot réserve un crédit par invité, puis settle les réussites et relâche le reste selon le résultat. Détails et limitations : `services/invitations/README.md` et le schéma Invitations.
""")

DOCUMENTS["01-architecture/README.md"] = doc("Guide des documents d’architecture", """
L’architecture repose sur le monorepo, le BFF, le Gateway, les services propriétaires, des bases séparées, une messagerie avec outbox et stockage objet privé. Diagrammes éditables et exports sont dans `diagrams/`. Chaque diagramme est un artefact autonome.

Les limites de disponibilité sont importantes : le Compose est local, plusieurs intégrations externes ne sont pas qualifiées, et `docs/E2E_VALIDATION.md` décrit une revue préproduction sans runtime.
""")

DOCUMENTS["02-data/ERD.md"] = doc("Modèle conceptuel des données", """
Les données sont réparties en bases propriétaires; les clés d’identité externes sont stockées comme `ownerSubject`/références UUID et ne sont pas des relations SQL inter-services. Les domaines majeurs comprennent profil/suppression, événements/cérémonies/agence, invités/groupes/accès, seating, designs/versions, media, invitations/RSVP/check-in/batches, Billing, Payments/partners, Wallet ledger/reservations, notifications, audit et jobs AI.

Le diagramme `diagrams/sources/erd-high-level.mmd` est volontairement de haut niveau et ne prétend pas à un FK cross-database. Les champs exacts sont dans `DATA_DICTIONARY.md` généré à partir des schémas Prisma et dans chaque `services/<domain>/prisma/schema.prisma`.
""")

DOCUMENTS["02-data/DATA_LIFECYCLE.md"] = doc("Cycle de vie des données", """
| Données | Création / mouvement | Suppression / durée |
|---|---|---|
| Compte et profil | Identité Keycloak; profil séparé dans Profile DB | demande de suppression existe; effacement réel multi-service doit être vérifié |
| Invités/imports | service Guests; job d’import possède statut et compteurs | service d’expiration d’import présent; durée exacte provient de config/code |
| Médias | quarantaine MinIO → vérification ClamAV/transcodage → bucket ready privé | suppression/tombstone et cleanup présents; politique légale doit être approuvée |
| Invitations | snapshots par version et batch en DB; PDF/ZIP en MinIO privé | expiration ZIP/cleanup configurés; rétention globale non approuvée |
| Paiements/ledger | snapshot ordre et transitions Payments; entrée Wallet append-only | conservation financière à définir par conseil juridique; aucune purge arbitraire |
| Audit/notifications | événements filtrés/dédupliqués et préférences | TTL/rétention à confirmer selon politique |
| Sessions | tokens chiffrés dans Redis; cookie navigateur opaque | TTL session 8 h; actualité à confirmer au code |

Les détails définitifs doivent refléter les migrations et paramètres de déploiement actifs, non une supposition documentaire.
""", "PARTIAL")

DOCUMENTS["02-data/DATA_CLASSIFICATION.md"] = doc("Classification des données", """
| Classe | Exemples du schéma / système | Contrôle attendu présent ou nécessaire |
|---|---|---|
| Identifiants/secrets | mots de passe, tokens OIDC, clés HMAC, SMTP, clés MinIO, tokens fournisseur | Secret hors Git; ne jamais journaliser; rotation et moindre privilège |
| Personnel invité | noms, emails/téléphones si fournis, notes | owner scope, réponse publique minimisée, éviter logs |
| Opérationnel privé | designs, listes d’invités, plans de salle, snapshots d’invitation | bases isolées, authorization, stockage privé |
| Transactionnel | références, montants, monnaie, statut provider, ledger crédits | idempotence, audit, restrictions mutation; données financières protégées |
| Audit | événement et métadonnées filtrés | append-only, allowlist et rôle support |
| Public contrôlé | invitation publiée via token signé | minimisation; ne pas exposer champs contact/notes |

Cette classification est un repère d’ingénierie et ne remplace pas une analyse juridique formelle ou une politique de conservation validée.
""")

DOCUMENTS["02-data/DATABASE_OWNERSHIP.md"] = doc("Propriété des bases", """
Chaque domaine actif a sa base Prisma et des credentials de service distincts; le provisioning est dans `infrastructure/postgres/` et `compose.yaml`. Domaines détectés : profile, events, guests, seating, designs, media, ai-design, billing, payments, wallet, notifications, invitations, audit, analytics et access. Les modèles générés ne doivent pas remplacer le schéma Prisma comme source éditable.

Pas de jointure SQL distribuée : les services échangent des identifiants et des événements/contrats. Cela impose des contrôles ownership dans chaque service et rend les suppressions cohérentes multi-service un sujet opérationnel.

La présence d’une migration n’indique pas son application dans une base existante. Avant un déploiement, inspecter `prisma migrate status` par service avec la DB correcte.
""")

DOCUMENTS["03-api/API_GUIDE.md"] = doc("Guide des API InvitaFlow", """
Le point d’entrée applicatif est Gateway (`apps/gateway`), généralement consommé par le BFF Web. Les routes domaine utilisent `/v1/...`; routes internes wallet/service tokens et routes publiques d’invitation ne suivent pas toutes le même schéma d’authentification.

Les requêtes authentifiées portent un token Keycloak access token; chaque service vérifie signature, issuer, audience et identité. Endpoints publics RSVP lisent un token d’invitation signé et ne délivrent pas d’identité utilisateur. Les actions mutantes financièrement/async requièrent une `Idempotency-Key` lorsqu’implémentée.

Les décorateurs routes sont dans les `*.controller.ts`; le Gateway possède un mapping explicite `apps/gateway/src/main.ts`. Aucune spécification OpenAPI centralisée n’a été repérée dans l’inventaire; le catalogue suivant est donc volontairement par familles plutôt que doublon d’un contrat généré.
""")

DOCUMENTS["03-api/API_CATALOG.md"] = doc("Catalogue des API", """
| Domaine | Route(s) observée(s) | Authentification | Usage |
|---|---|---|---|
| Profile | `GET/PUT /v1/me`, `/v1/me/deletion-request` | Bearer Keycloak, subject du token | profil et demande de suppression |
| Events | `GET/POST /v1/events`, `GET/PATCH /:eventId`, publish/cancel, ceremonies | Bearer + ownership | événement et cérémonies |
| Guests | `/v1/events/:eventId/guests`, groups, ceremonies, imports | Bearer + event ownership | invités, accès et import |
| Seating | `/v1/events/:eventId/ceremonies/:ceremonyId/seating` | Bearer + ownership | tables/placement/import |
| Designs | `/v1/events/:eventId/designs` et versions | Bearer + ownership | catalogue et design JSON |
| Invitations | `/v1/events/:eventId/invitations/batches`, download, cancel; `/v1/public/invitations/:token` | Bearer propriétaire / signature publique | génération, lecture et RSVP |
| Check-in | `/v1/events/:eventId/check-in/scan`, list | Bearer + rôle/ownership | présence par cérémonie |
| Wallet | `/v1/wallet/me`, transactions | Bearer audience `wallet-api` | soldes et ledger de crédits |
| Payments | `POST /v1/payments`, `GET /v1/payments/me`, `/v1/payments/:id`; callbacks provider | Bearer, clé d’idempotence; callback dépend contrat | checkout/consultation |
| Billing | `/v1/pricing`, `/v1/quotes`, finance-admin schedules | Bearer + rôle | catalogue, devis et planning |
| Media | `/api/assets` create/upload/complete/download/delete | Bearer + owner scope | asset lifecycle |
| Notifications | `/v1/notifications`, preferences | Bearer + owner scope | inbox/préférences |
| Audit | moderation reports/events | role support/admin; ingestion service-to-service | revue et audit |
| AI Design | `/v1/events/:eventId/designs/:designId/ai-jobs` | Bearer audience `ai-design-api` | jobs de proposition |

Ce tableau n’est pas une liste exhaustive des sous-routes. Vérifier les contrôleurs avant l’intégration et les contrats type de chaque handler.
""")

DOCUMENTS["03-api/ERROR_CATALOG.md"] = doc("Catalogue des erreurs API", """
Les services Nest utilisent exceptions HTTP explicites, notamment erreurs de validation `400`, absence/invalidité de bearer `401`, refus rôle/ownership `403` ou réponse masquée `404`, ressource absente `404`, conflit/idempotence `409`, limite/quota `429`, intégration upstream indisponible `503`. Les formes JSON exactes diffèrent entre handlers/proxies; ne pas supposer un envelope global uniforme sans lire le contrôleur.

Gateway retourne certaines erreurs normalisées d’indisponibilité (ex. `profile_service_unavailable`). Les erreurs fournisseurs paiement ne doivent pas exposer tokens/provider bodies. Pour observabilité, corréler via `x-request-id` et `x-correlation-id` lorsque le chemin le préserve.

Références: `apps/gateway/src/main.ts`, contrôleurs et guards `services/*/src/`.
""")

DOCUMENTS["03-api/EVENT_CATALOG.md"] = doc("Catalogue des événements métier", """
| Event type observé | Producteur | Consommateur / usage | Version / livraison |
|---|---|---|---|
| `payment.created.v1` | Payments | analytics/finance selon routage | outbox, version 1 |
| `payment.succeeded.v1` et `.v2` | Payments | Wallet crédit achats; autres projections | confirmé serveur, outbox transactionnelle |
| `payment.failed.v1` | Payments | projection/notifications selon routage | outbox |
| `payment.refunded.v1` et `.v2` | Payments | Wallet reversal pour v1 vérifié | outbox/idempotency |
| `payment.reconciliation-anomaly.v1` | Payments | supervision finance | outbox |
| `billing.price-schedule-published.v1` | Billing | projections/consommateurs configurés | outbox |
| `guest.created.v1`, `guest.updated.v1`, `guest.archived.v1` | Guests | notifications/analytics/audit selon binding | envelope versionnée |
| `guest.import.started/validated/completed.v1` | Guests | parcours async et analytics | outbox, job d’import |
| `seating.tables.imported.v1`, `seating.assignment.saved.v1` | Seating | notifications/analytics selon routing | outbox |
| `invitation.render.requested.v1` | Invitations | worker Rendering | queue durable, job DB reprenable |
| `invitation.rsvp.updated.v1` | Invitations | notifications organisateur | outbox |
| `media.background-removal.requested.v1` | Media | worker traitement image | outbox |
| `guest.import.expired.v1` | Guests | audit/projections selon binding | outbox |

L’enveloppe emploie `eventId`, `eventType`, `eventVersion`, `occurredAt`, `producer`, `correlationId`, `causationId`, `payload` dans plusieurs publishers. Le retry/attempts/DLX diffère par consumer; par exemple Notifications a limite de dix tentatives et DLX, Wallet indique retry 30 s puis dead queue. Le catalogue est le résultat des chaînes littérales actuelles, pas un contrat exhaustif. Binding exact et ordering doivent être revérifiés dans `infrastructure/rabbitmq/definitions.json` et consumers.
""", "PARTIAL")

DOCUMENTS["03-api/INTEGRATION_GUIDE.md"] = doc("Guide d’intégration", """
1. Passer par Gateway ou BFF documenté; ne pas exposer les endpoints service internes.
2. Obtenir un access token Keycloak adapté à l’audience du service et vérifier que le scope/client la fournit.
3. Pour les mutations async/financières, réutiliser une clé d’idempotence stable par opération logique.
4. Respecter les formats DTO et contrôles de taille du contrôleur propriétaire; ISO-8601 avec offset et IANA timezone pour Events.
5. Gérer pagination cursor si fournie; éviter retry aveugle non idempotent.
6. Pour les événements, consumer doit dédupliquer event ID, traiter une transaction localement et ack seulement après commit.
7. Utiliser les environnements Mock/Mailpit pour local. Aucun secret merchant ou SMTP dans navigateur, dépôt ou logs.

Contrats d’intégration externes FlexPay/EasyPay réels: NOT VERIFIED; demander la documentation fournisseur, ne pas extrapoler endpoints.
""")

DOCUMENTS["04-security/SECURITY_ARCHITECTURE.md"] = doc("Architecture de sécurité", """
## Identité et sessions

Keycloak est l’issuer OIDC. Le Web BFF utilise authorization code + PKCE S256, state/nonce, validation d’issuer/audience/azp et claim d’email vérifié. Redis conserve des tokens chiffrés; cookie opaque HttpOnly/SameSite et secure en production. Realm export active email verification, et l’email scope standard émet `email` et `email_verified`.

## Autorisation et frontières

Gateway exige un Bearer pour la plupart des proxys; les services refont validation et contrôle de propriétaire. Routes internes wallet ont `X-Service-Token` et idempotency; public invitation est scindé à un jeton signé minimisé. Les rôles Keycloak séparent client et admin. Le réseau Docker isole DB/broker/stockage, mais Compose n’est pas une politique réseau de production.

## Données et sécurité des fichiers

Médias écrivent d’abord en quarantaine; taille/MIME/magic bytes/dimensions/scanner ClamAV/décodage Sharp contrôlés avant publication prête. Bucket ready privé, URLs signées courte durée et scopes précis. PDF/ZIP invitations privés.

## Paiement et fraude

Montant/monnaie viennent du snapshot Billing côté serveur. Callback/webhook doit vérifier la preuve serveur et ses références avant transition/ledger. Idempotence, références uniques, outbox, ledger append-only et contrainte de solde sont des contrôles présents. Un fournisseur externe n’est pas qualifié sans contrat sandbox.

## Protections à vérifier

CSRF BFF, CSP, XSS dans rendu, SSRF, IDOR entre comptes, rate limiting, injection SQL/commandes, replay RSVP/QR, concurrence double check-in et paramètres proxy doivent rester dans plan de tests. Ne pas conclure « sécurisé » sur une simple revue statique. Liens : `THREAT_MODEL.md`, `SECURITY_TEST_PLAN.md`, guards, politiques MinIO, web CSP et payment providers.
""")

DOCUMENTS["04-security/THREAT_MODEL.md"] = doc("Modèle de menace STRIDE", """
| STRIDE / actif | Scénario | Contrôle observé | Risque résiduel / action |
|---|---|---|---|
| Spoofing / compte | voler code/session ou forger token | PKCE/state/nonce/JWKS; cookie opaque | tester fixation, replay, refresh et révocation en runtime |
| Tampering / event | utilisateur modifie ID d’un autre owner | ownerSubject guards | tests inter-comptes sur chaque domaine |
| Repudiation / paiement | contester callback ou crédit | outbox, références/idempotence, ledger append-only | audit complet et preuve marchand réel à valider |
| Information disclosure / invités | endpoint public fuite email/téléphone | projection publique minimisée, storage private | tests de réponses/routes directes et logs |
| Denial of service / média | fichier géant ou job IA sans fin | limites image, quotas AI, retries | charge/reverse proxy limits non mesurés |
| Elevation / admin | token mal audiencé ou rôle UI seul | guards Keycloak, audiences | tests rôle et endpoints directs |
| SSRF / image/provider | URL user atteinte depuis backend | politiques à examiner par feature | audit SSRF ciblé et allowlists |
| XSS / invitation | contenu événement injecté dans HTML/PDF | encodage template / snapshot | tests payload HTML/URL et rendu Chromium |
| Replay / QR/RSVP | token rejoué ou double scan | token HMAC, unicité de check-in par cérémonie | concurrency/API test; rotation HMAC invalide tous les liens |
| Malware / upload | fichier polyglotte/infecté | quarantine, clamd, decode/re-encode, private bucket | runtime test ClamAV et limites |
| Queue poison | message invalide retried | retry/DLX partiels | confirmer bindings et procédure DLQ par consumer |

Les contrôles et résiduels sont basés sur code/config observés; ce tableau ne constitue pas une certification ni une pentest.
""", "PARTIAL")

DOCUMENTS["04-security/AUTHORIZATION_MODEL.md"] = doc("Modèle d’autorisation", """
Identité principale: `sub` Keycloak comme owner subject. Token doit être signé par le realm attendu, non expiré et porter audience appropriée; un email valide n’est pas preuve de vérification, `email_verified=true` est requis côté Web. `CUSTOMER` est rôle initial annoncé dans realm export; rôles finance/support/admin/super-admin gouvernent opérations privilégiées.

Gateway ne remplace pas authorization métier. Chaque handler/service doit vérifier relation du sujet au `eventId`, asset, batch, design ou paiement. Les endpoints internes doivent rester privés réseau et exiger service token. Les routes publiques n’acceptent que les données publiées par token invitation.

Matrice détaillée: routes `services/*/src/*controller.ts`, guards `identity.guard.ts` et `internal.guard.ts`, rôle policies, Gateway allow routes. Test de non-accès A→B requis avant release.
""")

DOCUMENTS["04-security/SECRETS_MANAGEMENT.md"] = doc("Gestion des secrets", """
`.env.example` contient des valeurs placeholder. `.env` est ignoré par Git. Secrets à fournir hors dépôt : mots de passe DB, Redis, Keycloak admin, `AUTH_SESSION_SECRET`, secrets MinIO/service tokens, `INVITATION_LINK_SECRET`, SMTP app password, API token marchand, credentials AI/provider et observability exporters.

Local Mailpit: `MAIL_PROVIDER=mailpit`, aucun secret SMTP requis. Test réel/production: définir explicitement `MAIL_PROVIDER=smtp` et identifiants SMTP. Provider payment reste mock local jusqu’au contrat externe qualifié; tokens FlexPay/EasyPay ne sont jamais `NEXT_PUBLIC` et ne doivent pas être journalisés.

Production devrait injecter depuis un gestionnaire de secrets, restreindre lecture aux containers nécessaires, éviter secrets dans arguments/process/logs, documenter rotation et révocation, séparer credentials dev/stage/prod. Le dépôt Compose actuel fait passer variables aux conteneurs; un coffre externe prod est une exigence de déploiement, pas déclarée implémentée.
""")

DOCUMENTS["04-security/SECURITY_CONTROLS.md"] = doc("Registre des contrôles sécurité", """
| Contrôle | Source | État |
|---|---|---|
| PKCE S256/state/nonce et session BFF | auth-session, callback | code présent; runtime non vérifié |
| Realm email verification / duplicate prevention | realm export, provision-email | config présente |
| Verified-email claim required | callback web | implemented |
| AuthZ service/role | guards/controller décorateurs | code présent, couverture incomplète |
| DB isolation | Prisma schemas, Postgres provisioners, Compose | config présente |
| Image quarantine/clamav/decode | Media | code et tests; scanner runtime non vérifié |
| MinIO least privilege/private bucket | infrastructure/minio, policies | configuration; runtime IAM non vérifié |
| Payment idempotency/ledger | Payments/Wallet | code/tests par service |
| HTTP security headers / request context | Gateway, Web CSP | code présente; proxy déploiement à vérifier |
| Observability alert rules | `infrastructure/monitoring` | config présente; alert delivery non testé |
| Vulnerability / dependency scan | CI workflows if configured | voir workflows; ne pas supposer audit exhaustif |

Ce registre se met à jour à partir des configs et rapports réellement exécutés.
""")

DOCUMENTS["04-security/INCIDENT_RESPONSE.md"] = doc("Procédure de réponse incident", """
## Séquence initiale

1. Déclarer incident, heure, service, environnement et coordinateur; éviter recopier PII/secrets dans tickets.
2. Évaluer risque immédiat : authentification compromise, fuite invités, accès objet, paiement/ledger, disponibilité.
3. Contenir via rotation/révocation ciblée des credentials/tokens, désactivation de clé/lien, isolation du service ou arrêt du worker selon périmètre; préserver journaux et DB.
4. Capturer request/correlation IDs, event IDs, références payment, commit/config version, fenêtre temps; conserver intégrité et accès restreint.
5. Récupérer depuis backup documenté seulement après validation de l’instantané et isolement de la cible; ne pas réécrire ledger/audit.
6. Vérifier santé, migrations, login, ownership, stockage, ledger/outbox et erreurs; communiquer selon obligations réglementaires et légales approuvées.
7. Rédiger cause, portée, données touchées, actions préventives et propriétaire.

Les contacts officiels produit sont `fucushd098@gmail.com`, `+243973431495`. Les astreintes, délais d’escalade, obligations de notification et RTO/RPO n’apparaissent pas définis dans le dépôt : les approuver avant production.
""", "PLANNED")

DOCUMENTS["04-security/SECURITY_TEST_PLAN.md"] = doc("Plan de tests sécurité", """
Priorités avant release:

1. AuthN: state/nonce incorrect/absent/rejoué, PKCE verifier mauvais, issuer/audience/azp invalide, expiration, email absent/non vérifié, refresh concurrent, logout/revocation.
2. Authorization A/B: tester accès direct event, guest, seating, design, asset, batch, invitation, paiement, wallet et admin avec un autre subject/role.
3. BFF/CSRF: cookies Secure/HttpOnly/SameSite, Origin, cross-site writes, cache/no-store.
4. Validation: JSON inattendu, pagination abusive, import ambigu/volumineux, IDs malformés, prompt, SVG/HTML, URL distantes.
5. Fichiers/SSRF: polyglot, archive expansion, MIME spoof, ClamAV timeout/infected, URL callback, bucket isolation.
6. Paiement: forged success, mauvais montant/devise/référence, duplicated/out-of-order IPN, concurrent browser-success/IPN, replay, mismatch et secret logs.
7. QR/RSVP/check-in: token altéré/expiré/révoqué, double RSVP/check-in concurrent, cérémonie non invitée.
8. Abuse/rate: login, email reset, création jobs/batches, upload, public endpoints.

Enregistrer environment/version, fixtures fictives, preuve/ID, expected/actual, résultat PASS/FAIL/BLOCKED/NOT TESTED. Ne pas faire de fuzz sur production sans autorisation.
""", "PLANNED")

DOCUMENTS["05-payments/PAYMENT_PROVIDER_ARCHITECTURE.md"] = doc("Abstraction des fournisseurs de paiement", """
Payments dépend de l’interface `PaymentProvider`, possède les ordres/attempts et orchestre les transitions. `MockPaymentProvider` est le chemin de développement. `FlexPayProvider` est présent mais README indique que contrat marchand et sandbox ne sont pas disponibles; il refuse volontairement checkout/vérification/refund. Des traces CinetPay historiques restent dans `provider` sans adapter vérificateur actif.

Ne pas conclure qu’EasyPay est prêt parce que des docs/variables de phase précédente existent : comparer `docs/PAYMENTS_EASYPAY.md` au code provider actuel. Aucun appel réel n’est attesté par les tests provider qui mockent fetch.

Un retour frontend ne crée pas de crédit. La preuve serveur met à jour le paiement et outbox atomiquement; Wallet traite les événements avec clé de ledger idempotente. Remboursement réel/settlement exige vérification fournisseur et politique finance.
""", "PARTIAL")

DOCUMENTS["05-payments/EASYPAY_INTEGRATION.md"] = doc("Intégration EasyPay état actuel", """
Le code et notes de phase antérieurs peuvent contenir une intention EasyPay. La source opérationnelle la plus actuelle à confirmer reste `services/payments/README.md` et le code `services/payments/src/providers/`: FlexPay est le seul adapter externe nommé actif mais non fonctionnel faute de contrat; Mock reste le fournisseur de test. Aucune route/structure EasyPay/IPN, méthode d’auth ou mapping de statuts ne doit être inventé.

Avant activation: obtenir documentation officielle RDC, CID/token sandbox, formats/headers, statut final checking-status, preuve IPN, montants/monnaies, délais, refunds et cas d’erreur; valider via serveur; masquer secrets; tester duplicate/out-of-order/mismatch; puis seulement activer une sélection de fournisseur explicitement documentée.

**État: NOT VERIFIED / BLOCKED_PROVIDER_CONTRACT.** Voir aussi `docs/PAYMENTS_EASYPAY.md` pour l’historique de conception.
""", "NOT VERIFIED")

DOCUMENTS["05-payments/PAYMENT_STATE_MACHINE.md"] = doc("États des ordres et paiements", """
Enums actuels `PaymentStatus`: `CREATED`, `PENDING`, `PROCESSING`, `SUCCEEDED`, `FAILED`, `CANCELLED`, `EXPIRED`, `REFUND_PENDING`, `REFUNDED`. `OrderStatus`: `CREATED`, `PAID`, `CANCELLED`, `REFUNDED`. Ce n’est pas un workflow complet et séquence exacte doit être prise du Payments service.

Diagramme : `diagrams/sources/payment-state-machine.mmd`. Un succès ne doit venir qu’après confirmation provider; le diagramme ne remplace ni invariants ni transaction.
""")

DOCUMENTS["05-payments/PAYMENT_IDEMPOTENCY.md"] = doc("Idempotence et réconciliation paiement", """
Le `POST /v1/payments` prend une `Idempotency-Key`; ordre/paiement enregistrent les refs, tentatives et snapshot catalogue. Un compare-and-set évite plusieurs initialisations concurrentes. Références transaction fournisseur et reçus webhook ont contraintes uniques. Une confirmation serveur, changement ordre et outbox sont dans une transaction DB. Wallet applique un idempotency key payment-scoped.

Pour exploiter: préserver même clé sur retry client d’une seule action, distinguer nouvel achat d’un retry, ne jamais considérer return URL comme vérification, garder mismatch en anomalie et vérifier DLQ. Voir `services/payments/README.md`, `payments.service.ts`, `wallet/payment-consumer.ts`.
""")

DOCUMENTS["05-payments/WALLET_ARCHITECTURE.md"] = doc("Architecture du Wallet de crédits", """
Wallet détient `available` et `reserved` unités de crédits et un ledger append-only. Le solde est une projection mise à jour avec l’entrée ledger dans une transaction; contrainte DB empêche solde négatif; trigger interdit update/delete ledger. Correction via reversal. Les réservations génération invitation sont consommées/libérées/settled idempotemment.

Crédits = unités d’usage. Ce n’est ni devise ni argent; ne pas afficher le wallet comme compte de paiement. L’acquisition peut dépendre d’un ordre `CREDIT_PURCHASE`, mais valeur monétaire vit dans Payment/Billing et le fournisseur.
""")

DOCUMENTS["05-payments/PAYMENT_VALIDATION_REPORT.md"] = doc("Rapport validation paiements", """
| Vérification | Résultat actuel |
|---|---|
| Mock provider unitaire | tests provider présents; exécution complète non attestée dans cette génération |
| Provider externe | FlexPay adapter volontairement indisponible faute contrat |
| EasyPay sandbox | NOT TESTED / credentials/contract not confirmed |
| IPN/reconciliation réel | NOT TESTED |
| Wallet crédit runtime | NOT TESTED |
| Idempotence par code | validations et contraintes présentes; simultanéité runtime requise |
| Montant/monnaie invalides | tests adapter existent par sous-domaines; intégration E2E requise |
| Secret handling | variables serveur; aucun secret à inclure dans docs/CI output |

Source historique `docs/E2E_VALIDATION.md` donne l’état runtime de sa date. Refaire ce rapport à chaque release avec logs anonymisés et IDs références faux.
""", "NOT VERIFIED")

DOCUMENTS["06-auth-email/KEYCLOAK_ARCHITECTURE.md"] = doc("Architecture Keycloak", """
Le realm export `invitaflow-dev` autorise inscription, login email, vérification et reset password, interdit emails dupliqués, sélectionne le login theme `invitaflow`, email theme `invitaflow`, et configure client public Web avec standard flow et PKCE S256. Client `invitaflow-web` a scopes `profile`, `email`, audiences API et rôles.

Le thème d’inscription montre l’acceptation légale; un FormAction serveur valide et enregistre les versions/time. `provision-legal-registration.mjs` lie le flow copié au realm; `provision-email.mjs` ajoute config SMTP/Mailpit et rend l’email requis. Les realms persistants ne sont pas écrasés par import : les jobs bootstrap configurent le realm au runtime Compose.

Secrets admin/SMTP restent server side. Aucun secret n’est décrit ici.
""")

DOCUMENTS["06-auth-email/REGISTRATION_FLOW.md"] = doc("Parcours inscription", """
1. L’utilisateur choisit inscription Keycloak.
2. Le profil exige email valide et champs d’identité/password gérés par User Profile.
3. La case légale n’est pas précochée; liens `/legal/cgu` et `/legal/confidentialite` s’ouvrent séparément.
4. FormAction légal serveur refuse si case absente; après création utilisateur, persiste `termsVersion`, `privacyVersion`, `termsAcceptedAt`, `privacyAcceptedAt` depuis manifest/code.
5. Realm `verifyEmail` et action `VERIFY_EMAIL` conduisent à l’email vérification.
6. Session applicative n’est créée qu’au callback OIDC après claim email vérifié.

Diagrammes séparés `registration-legal-acceptance` et `email-verification`. Un walkthrough Keycloak reste à exécuter après build/runtime.
""")

DOCUMENTS["06-auth-email/LEGAL_ACCEPTANCE_FLOW.md"] = doc("Parcours acceptation légale", """
Le contrat des versions est centralisé dans `packages/legal-contract/legal-manifest.json`. La page d’inscription expose les liens CGU/Confidentialité et contrôle client pour l’UX; la sécurité vient du FormAction provider qui exige la valeur `accepted` côté serveur. Le handler attache timestamps UTC et versions à l’utilisateur.

Attributs : `termsVersion`, `privacyVersion`, `termsAcceptedAt`, `privacyAcceptedAt`. Case unchecked par défaut; l’utilisateur doit cocher activement pour chaque nouvelle inscription. Les versions doivent être mises à jour selon le manifest et stratégie de migration.

Les documents légaux eux-mêmes sont des sources dans `apps/web/content/legal/`; leur conformité et exactitude nécessitent validation juridique par l’opérateur.
""")

DOCUMENTS["06-auth-email/EMAIL_VERIFICATION_FLOW.md"] = doc("Vérification email", """
À la création, Keycloak ne doit pas déclarer l’adresse vérifiée seulement parce qu’elle passe le format : `emailVerified` reste faux et action `VERIFY_EMAIL` doit être requise lorsque `verifyEmail=true`. Keycloak envoie un lien one-time signé selon sa config d’expiration. Le lien vérifié met le compte à `emailVerified=true`. À l’authentification suivante, standard scope `email` mappe `email` et `emailVerified` vers claims ID token `email`, `email_verified`; le callback exige les deux.

Le mail local utilise Mailpit; test courrier réel/production nécessite `MAIL_PROVIDER=smtp`. Runtime séquentiel registration→email→verification→OIDC doit être validé. Diagramme `email-verification-sequence.mmd`.
""")

DOCUMENTS["06-auth-email/SMTP_CONFIGURATION.md"] = doc("Configuration SMTP", """
| Mode | Paramétrage |
|---|---|
| Local | `MAIL_PROVIDER=mailpit`; SMTP interne `mailpit:1025`, aucune auth |
| Test avec email réel | `MAIL_PROVIDER=smtp`; identifiants SMTP dans secrets environment |
| Production | `MAIL_PROVIDER=smtp`; credentials et host/from validés par opérateur |

Variables: `KEYCLOAK_SMTP_HOST`, `KEYCLOAK_SMTP_PORT`, `KEYCLOAK_SMTP_FROM`, `KEYCLOAK_SMTP_FROM_DISPLAY_NAME`, `KEYCLOAK_SMTP_USER`, `KEYCLOAK_SMTP_PASSWORD`, `KEYCLOAK_SMTP_STARTTLS`, `KEYCLOAK_SMTP_SSL`. Pour l’option Gmail indiquée au dépôt : `smtp.gmail.com`, port 587, STARTTLS; mot de passe d’application, jamais mot de passe normal. `.env.example` laisse secret vide; `.env` est ignoré. Aucun mot de passe n’est consigné dans les logs.

Le service provision email valide mode/credential et applique réglages au realm. Ne jamais mettre secret dans JSON realm, front, documentation ou output console.
""")

DOCUMENTS["07-testing/TEST_STRATEGY.md"] = doc("Stratégie de test", """
## Niveaux

- Unit: services, parseurs, règles état/idempotence, composants purs et package design-document.
- Integration: PostgreSQL/RabbitMQ/MinIO/Keycloak/proxy à exécuter contre stacks isolées et services réels simulés ou test.
- Contract: API DTO, event envelopes, audience OIDC, signatures provider/MinIO, interop snapshot renderer.
- E2E navigateur: création profil/event/import/seat/design/batch/RSVP/check-in; auth email; admin/agency/partner.
- Security: plan `04-security/SECURITY_TEST_PLAN.md`, A/B access, IPN, files, QR, web sessions.
- Performance/accessibility/mobile: plans dédiés; aucune mesure PASS sans rapport récent.
- Payment sandbox: bloqué jusqu’au contrat et credentials autorisés.

## Répertoire de preuves

`services/*/*.spec.*`, `packages/design-document/test/`, `apps/web/src/**/*.test.mjs`, `load-tests/k6/`. `pnpm test` utilise Turbo plus tests racine; services avec scripts propres. Le CI Node version/matrix est décrit dans README/workflow source.

## Traceabilité indicative

| Requis | Famille tests |
|---|---|
| FR-AUTH | `apps/web/src/lib/email-verification.test.mjs`, session refresh; Keycloak runtime manquant |
| FR-EVENT | Events date/program specs |
| FR-GUEST | import parser/list specs |
| FR-SEAT | capacity/guests specs |
| FR-DESIGN / FR-AI | DesignDocument/Designs/AI tests |
| FR-INV | Invitations token/layout/image/service specs |
| FR-MEDIA | validation, antivirus, transcode, presign specs |
| FR-PAY / Wallet | payments provider/checkout/list, wallet ledger/routing specs |

Un test unitaire ne prouve pas un parcours connecté E2E.
""")

DOCUMENTS["07-testing/E2E_VALIDATION.md"] = doc("Validation end-to-end", """
État de référence : `docs/E2E_VALIDATION.md` historique. Sa matrice signale runtime Docker indisponible, EasyPay sandbox absent et divers E2E NOT TESTED. Ce dossier d’ingénierie doit conserver la preuve (date, commit, runtime, comptes fictifs, logs sans PII, résultat) avant de reclasser une fonction.

Pour une campagne: créer deux comptes A/B, deux événements, importer petit XLSX fictif, placer invités, choisir design, rendre une invitation, RSVP, tester QR/check-in, valider wallet/pay provider sandbox seulement si activé, et refaire un accès croisé A→B. Ne pas réutiliser données client réelle.
""", "NOT VERIFIED")

DOCUMENTS["07-testing/SECURITY_TEST_REPORT.md"] = doc("Rapport de tests sécurité", """
Aucun rapport de pentest indépendant ou campagne complète actuelle trouvée. Unit tests ciblés des policies, auth helper, upload, payment parsing et idempotence existent selon répertoires, mais ne valent pas rapport de sécurité complet.

Résultat de référence: **NOT TESTED** pour DAST, SAST approfondi, exploit inter-user E2E, payment-IPN sandbox et contrôle de secrets de déploiement. Remplir à l’exécution avec outils/version, commit, périmètre, résultats, exceptions, preuves et plan de correction.
""", "NOT VERIFIED")

DOCUMENTS["07-testing/PERFORMANCE_TEST_PLAN.md"] = doc("Plan de performance", """
Les scripts k6 existants couvrent RSVP, lots d’invitations, guest search et check-in (`load-tests/k6`). Avant mesure, utiliser base et identité de test, purger fixtures, vérifier quotas et coûts. Cibles à approuver: p50/p95/p99, débit, concurrence et erreur par endpoint; aucun seuil SLA validé n’est défini dans le dépôt.

Scénarios: liste invités et curseurs, recherche guest, RSVP public, check-in concurrent, lots de rendu, upload/scanner, jobs IA avec quota, checkout/init/provider mock, retries queue. Mesurer CPU/mémoire DB/Redis/Rabbit/Chromium, queue age/DLQ et objet MinIO.
""", "PLANNED")

DOCUMENTS["07-testing/PERFORMANCE_REPORT.md"] = doc("Rapport performance", """
Aucun rapport de benchmark reproductible, environnement/capacité/objectifs approuvés détecté. Statut **NOT GENERATED**. Exécuter le plan après fixation des cibles, noter données synthétiques, config, commit, courbe de charge, saturation, résultats et limites. Ne pas extrapoler un test local en capacité de production.
""", "NOT VERIFIED")

DOCUMENTS["07-testing/MOBILE_VALIDATION.md"] = doc("Validation mobile", """
Statut **NOT TESTED** dans cette campagne documentaire. Capturer viewport réel ou émulateur avec taille OS/navigateur, scénario, orientation, safe-area, états clair/sombre, zoom clavier, upload/camera, scan BarcodeDetector fallback et téléchargements. Guides captures: `10-user-guides/screenshots/end-user/SCREENSHOT_CHECKLIST.md`.
""", "NOT VERIFIED")

DOCUMENTS["07-testing/ACCESSIBILITY_VALIDATION.md"] = doc("Validation accessibilité", """
Composants visuels comportent des labels/focus/thèmes; aucune conformité WCAG globale certifiée trouvée. Audit à mener au clavier et lecteur d’écran: noms accessibles, ordre/focus, contrastes normal/grand texte, erreurs liées aux champs, contrastes dark mode, zoom 200/400 %, targets tactiles, langue, tableaux et graphiques.

Documenter version outil/navigateur, pages, nombre violations critiques, contrast ratio et exceptions. Statut **NOT VERIFIED**.
""", "NOT VERIFIED")

DOCUMENTS["08-operations/RUNBOOK.md"] = doc("Runbook local", """
## Pré-requis et démarrage

Node 24–26, pnpm 12.8, Docker Desktop Compose v2. Copier `.env.example` vers `.env`, générer des secrets uniques pour local, conserver `MAIL_PROVIDER=mailpit`, installer `pnpm install`, puis `pnpm infra:up`. Compose est une stack dev uniquement.

```powershell
docker compose config --quiet
docker compose up -d
docker compose ps
docker compose logs --tail 200 <service>
docker compose restart <service>
docker compose build <service>
docker compose down
```

Ne pas utiliser `down -v` sans plan explicite de destruction données.

## Diagnostic

1. `docker compose ps --all`, regarder health/exit code.
2. `docker compose logs --tail 200 keycloak keycloak-email-init keycloak-legal-flow-init` pour bootstrap.
3. Inspecter `postgres` et `*-db-init` sans afficher passwords; contrôler migrations par service.
4. Vérifier Ready `/health/ready`, Gateway `/health/live`, Mailpit UI localhost:8025, RabbitMQ management localhost:15672, MinIO console localhost:9001.
5. Pour queue: dashboard RabbitMQ, backlog/retry/DLQ; ne pas supprimer messages sans capture.
6. Pour media: état d’asset, clamd health, quarantine/ready buckets; ne jamais rendre bucket public.
7. Pour paiement: commencer `PAYMENT_PROVIDER=mock`; ne jamais configurer provider réel sans contrat/sandbox confirmé.
8. Exportes observability via `docker compose --profile observability up -d` et credentials uniques.

## Migrations / bootstrap

Entrypoints DB-init créent roles/bases sur volume neuf/existant; Keycloak bootstrap audiences/legal/email est idempotent; MinIO bootstrap policies/accounts. Vérifier logs job exit 0 et migration status. Ne pas reset volumes pour corriger un bootstrap.

## Incidents fréquents

- Mail non reçu local : confirmer `MAIL_PROVIDER=mailpit`, job `keycloak-email-init`, mailpit health et boîte UI 8025.
- SMTP échoue : confirmer mode explicit `smtp`, host/port/TLS/username/password depuis secret sans imprimer sa valeur.
- Callback OIDC refusé : regarder cause/ID correlation sans afficher token; vérifier redirect URI, realm/client, PKCE state et `email_verified`.
- Rendu attente : examiner outbox, queue, worker, batch items, MinIO et DLQ; retry limité selon service.
- Docker unavailable : vérifier Desktop engine/socket puis `docker info`; ne pas déclarer health inconnue PASS.

Contacts opérateur: `fucushd098@gmail.com`, `+243973431495`.
""")

DOCUMENTS["08-operations/DEPLOYMENT.md"] = doc("Guide déploiement", """
Compose décrit le développement local. Un déploiement réel doit dériver une configuration dédiée, TLS/proxy, secret store, DB/broker/object store managés ou durcis, sauvegardes externes, réseau minimal, image pinning/SBOM, permissions, alerting et release progressive. Aucune topologie production HA complète n’est déclarée prête ici.

Checklist: revue diff/migrations, tests+build CI, variables requises par service sans valeurs en Git, nouveau realm/provider/bootstrap, SMTP `smtp` explicit, clés HMAC et audiences, migrations dans l’ordre contrôlé, readiness et smoke tests, alertes/backup, plan rollback. Ne pas exporter port de DB/broker vers public. Aucun EasyPay token/client ID dans navigateur.
""", "PARTIAL")

DOCUMENTS["08-operations/MIGRATIONS.md"] = doc("Gestion des migrations", """
Chaque service DB a `prisma/schema.prisma`, config et historique de migrations propre. Les migrations sont indépendantes; identité Keycloak a son propre DB/realm. Nouvelle image ne prouve pas que migration live est appliquée.

Avant déploiement: sauvegarde, `prisma migrate status` par service avec URL correcte, revue SQL et ordre dépendances; appliquer `migrate deploy` selon runbook contrôlé. Aucun `migrate reset` sur volume utilisateur. Vérifier version migration, health et données test.

Les migrations réellement présentes sont sous `services/<name>/prisma/migrations/`. Base vivante actuelle: NOT VERIFIED.
""", "PARTIAL")

DOCUMENTS["08-operations/BACKUP_RECOVERY.md"] = doc("Sauvegarde et restauration", """
Référence implémentée : `docs/OPERATIONS_BACKUPS.md`, `scripts/backup-postgres.mjs`, `scripts/restore-postgres.mjs`. pg_dumpall inclut toutes bases/rôles et données sensibles; archive chiffrée, restreinte et hors Git. Objets MinIO ne sont pas couverts par dump PostgreSQL: les répliquer dans emplacement indépendant avec credentials least-privilege et checksum/versions.

Ordre restore proposé: créer cible isolée vide → restaurer PostgreSQL cluster/roles → aligner migrations → restaurer buckets au même point de reprise → config/secrets compatibles → bootstrap sans remplacer données → smoke auth/owner/event/wallet/invitation/object checks → basculer trafic contrôlé.

Drill réel NOT TESTED; RPO/RTO et durée de rétention non approuvés.
""")

DOCUMENTS["08-operations/DISASTER_RECOVERY.md"] = doc("Plan de reprise après sinistre", """
| Scénario | Priorité reprise proposée | Dépendances |
|---|---|---|
| Perte DB service | restaurer DB/roles, appliquer schéma compatible | backups PostgreSQL, migrations, secrets |
| Perte MinIO/objets | restaurer bucket au même recovery point | réplication objets, policies/accounts, keys |
| Keycloak indisponible | restaurer DB realm puis valider OIDC/keys/required actions | backup DB et secrets issuer |
| RabbitMQ panne | restaurer définitions puis requeue depuis outbox | DB outboxes, DLX/retry config |
| Secret compromis | révoquer/rotater clé ciblée, invalidation éventuelle sessions/tokens | procédure provider, opérateur |
| Paiement mismatch | geler crédits/payout, rapprocher merchant | logs minimisés, order/ref, contrat |

RTO/RPO: **PROPOSED, à approuver**. Chiffrer backup en transit/repos; tester restauration isolée avant failback, vérifier balances ledger, invitation snapshots et objets référencés. Consigner communications clients/autorités selon avis juridique. Aucune promesse de continuité active-active n’est faite.
""", "PLANNED")

DOCUMENTS["08-operations/OBSERVABILITY.md"] = doc("Observabilité", """
Actuel: health liveness/readiness dans services, request/correlation ID et Prometheus HTTP metrics Gateway, Prometheus alert rules, RabbitMQ metrics plugin/export, exporters PostgreSQL/Redis dans profil `observability`, Grafana, OTEL Collector, Tempo/Loki configurés partiellement. `infrastructure/monitoring/README.md` précise que host CPU/memory/disk et business/payment metrics sont incomplets.

Pas de trace distribuée end-to-end prouvée. Pas d’alert delivery/runtime collector qualifié. Dashboards/alerts doivent être testés avec cible inaccessible et backlog/DLQ.

Diagramme `observability-architecture.mmd`; configuration `infrastructure/monitoring/` et RabbitMQ definitions.
""", "PARTIAL")

DOCUMENTS["08-operations/LOGGING.md"] = doc("Politique de journalisation", """
Gateway propage `x-request-id`, `x-correlation-id`; logs doivent aider le diagnostic sans bearer, cookie, mot de passe, SMTP secret, provider token, invitation URL/QR full token, email/téléphone ou texte libre d’invité. Audit event utilise allowlist de métadonnées et évite PII/free text selon README racine.

Uniformiser niveaux, champs timestamp/service/environment/errorCode, IDs corrélation et redaction. Définir rétention/accès/immutabilité du backend logs pour le déploiement; ces durées ne sont pas garanties dans le repo.
""")

DOCUMENTS["08-operations/MONITORING_ALERTS.md"] = doc("Matrice alertes", """
| Condition | Sévérité proposée | Signal disponible | Action |
|---|---|---|---|
| Gateway/DB/Redis exporter absent | SEV2 | Prometheus target alert | valider readiness/config et credentials |
| PostgreSQL connexions saturées | SEV2 | postgres exporter/alert rules | examiner pool et saturation avant scale |
| RabbitMQ backlog durable | SEV2/SEV3 | queue ready/unacked | examiner consumer/outbox et âge |
| Dead-letter non vide | SEV2 | queue DLX | capturer messages non PII, traiter cause, replay idempotent |
| Payments reconciliation anomaly | SEV1 finance selon exposition | DB issue/outbox | geler crédit/refund/payout, rapprocher preuve provider |
| Upload scanner failure | SEV2 sécurité | service health/error | ne pas marquer média ready; rétablir scanner |
| MinIO object mismatch | SEV1/2 | inventory/audit service | conserver échec fermé et restaurer/manually review |

Sévérités/action proposées car seuils et escalade ne sont pas validés. Signaux et règles effectives: `infrastructure/monitoring/alerts.yml`.
""", "PARTIAL")

DOCUMENTS["08-operations/TROUBLESHOOTING.md"] = doc("Dépannage InvitaFlow", """
- **Compose démarre mais Web attend** : consulter jobs Keycloak/legal/email et DB-init, puis health des services dépendants; ne pas contourner une migration.
- **SMTP dev** : Mailpit par défaut; retrouver message sur port UI 8025. SMTP externe seulement après `MAIL_PROVIDER=smtp` explicite et credentials injectés.
- **Email login/session échoue** : audience email scope, `email`/`email_verified`, redirect URI, `email_verified=true`, state/PKCE et Keycloak event log sans tokens.
- **Import invités** : vérifier format colonne/mapping, job preview/commit et limites; ne pas réimporter un commit sans idempotency.
- **PDF pas prêt** : regarder InvitationBatch/BatchItem status, outbox `invitation.render.requested.v1`, queue worker, retries, stockage MinIO; télécharger seulement completed.
- **Wallet/credits** : consulter ledger et reservation ref, paiement success verified et idempotency. Ne jamais corriger solde en SQL.
- **Media non ready** : asset reste en quarantine tant que ClamAV/validation/transcode/publication inachevée.
- **Payment return success sans crédit** : retour navigateur ne crédite pas; inspecter statut serveur/provider, outbox et consumer Wallet.
""")

DOCUMENTS["09-release/RELEASE_CHECKLIST.md"] = doc("Liste de contrôle de release", """
- [ ] Scope et migrations relus par domaine; aucune migration destructive non approuvée.
- [ ] Secrets absents du Git; config correcte par environnement, SMTP explicitement smtp en production.
- [ ] CI lint/typecheck/tests/build verts pour commit taggé; versions images fixées.
- [ ] Bootstrap Keycloak/legal/email idempotent et scopes claims validés runtime.
- [ ] DB migration status et backup pris; procédure restore prête.
- [ ] Smoke auth verified email, user A/B ownership, event/import/design/render/pdf, RSVP/check-in.
- [ ] Paiement sandbox de contrat officiel ou provider Mock retiré de la production; webhooks idempotents/montant vérifié.
- [ ] Media ClamAV/MinIO buckets privés, object refs vérifiés.
- [ ] Alert rules et DLQ testées, contacts/escalade convenus.
- [ ] Mobile/clavier/contraste, guide capture réel et legal pages vérifiés.
- [ ] Rollback/communications et propriétaires confirmés.

Cocher avec preuve horodatée, pas par déclaration.
""", "PLANNED")

DOCUMENTS["09-release/RELEASE_READINESS_REPORT.md"] = doc("Rapport de préparation release", """
**Conclusion : BLOCKED pour une déclaration de production prête** tant que les parcours runtime, paiement provider et restauration ne sont pas vérifiés.

| Domaine | Statut |
|---|---|
| Infrastructure | PARTIAL — Compose documenté, runtime dépendant |
| Auth / legal | PARTIAL — code config présent, walkthrough runtime requis |
| Email | PARTIAL — Mailpit/SMTP provision configurés; délivrabilité non testée |
| Events / guests / import / seating / design | PARTIAL — routes/tests présents, E2E cross-owner non attesté |
| Rendering / PDF / QR / RSVP / check-in | PARTIAL — service/tests/config présents, runtime non vérifié |
| Wallet | PARTIAL — ledger et tests présents, flux réel non testé |
| Payments / EasyPay | BLOCKED/NOT VERIFIED — fournisseur externe non qualifié |
| Storage / media | PARTIAL — code scanner/policies présents; runtime requis |
| Agency / partner | PARTIAL — schémas/routes; permissions/payout non E2E |
| Mobile / accessibility | NOT TESTED |
| Security | PARTIAL — contrôles code, tests adversariaux runtime manquants |
| Backups | BLOCKED — restore drill non effectué |
| Observability | PARTIAL — configuration, collecte/alerte réelle non vérifiée |
| Deployment / rollback | BLOCKED — stack local, runbook production final absent |

Sources : `docs/E2E_VALIDATION.md`, README services, CI et code actuel. La date de source de l’E2E peut être plus ancienne; revalider avant usage release.
""", "BLOCKED")

DOCUMENTS["09-release/PRODUCTION_READINESS.md"] = doc("Critères de préparation production", """
Passage production requiert des environnements isolés, TLS, secret store, SMTP réel contrôlé, fournisseur paiement officiel contractuel, policy legal/data validée, migrations vérifiées, backups et restore drill, tests owner isolation, monitoring/alert route, rate limiting et sécurité reverse proxy, capacité mesurée, rollback éprouvé, support/escalade.

Ces critères sont une liste actionnable, non l’affirmation que tous les mécanismes existent déjà. Appuyer chaque PASS sur commit+run ID+evidence.
""", "PARTIAL")

DOCUMENTS["09-release/ROLLBACK_PLAN.md"] = doc("Plan de retour arrière", """
1. Détecter régression et arrêter l’extension trafic; préserver IDs/error/metrics et n’interrompre pas transaction financière au milieu sans vérifier.
2. Geler workers/consumers seulement selon runbook, capturer outbox/queue/DB state, empêcher doubles effets.
3. Revenir à une image précédente uniquement si schema backward-compatible; sinon appliquer procédure forward fix ou restore isolé approuvé.
4. Ne jamais `down -v`, reset DB, delete ledger/outbox, rejouer DLQ non idempotent.
5. Vérifier checkout orders, wallet invariants, access sessions, renders/object links puis rouvrir workers.
6. Communiquer état/impact; consigner décisions et preuve.

Migrations rollback spécifique n’est pas garanti par Prisma; élaborer migration corrective testée pour chaque release. Contact opérateur: email/téléphone au README.
""")

DOCUMENTS["09-release/CHANGELOG.md"] = doc("Journal de changement documentation", """
## 1.0 - 2026-10-07

- Première documentation structurée produit, architecture, données, APIs, sécurité, paiements, auth/email, opérations, tests, releases et guides.
- 32 sources diagrammes Mermaid et exports SVG/PNG générés.
- 16 exports Word et outillage local de rendu/validation ajoutés.
- Captures applicatives non générées, faute de runtime authentifié démontré.
""")

DOCUMENTS["10-user-guides/END_USER_GUIDE.md"] = doc("Guide visuel utilisateur InvitaFlow", """
Ce guide accompagne l’organisateur francophone. Les captures demandées sont **NOT GENERATED** tant qu’une application réelle accessible avec compte de démonstration documenté n’est pas disponible. Ne pas remplacer par maquettes fictives. Liste: `screenshots/end-user/SCREENSHOT_CHECKLIST.md`.

## 1. Accéder et créer un compte

Ouvrez l’adresse fournie par InvitaFlow, choisissez création de compte, indiquez une adresse email que vous pouvez ouvrir, créez un mot de passe conforme, lisez puis cochez les Conditions Générales d’Utilisation et la Politique de confidentialité. La case est obligatoire.

## 2. Vérifier et se connecter

Ouvrez l’email de vérification et suivez son lien avant expiration. Si vous ne voyez pas le message, vérifiez Courrier indésirable ou demandez de l’aide à `fucushd098@gmail.com`, `+243973431495`. Connectez-vous ensuite avec le compte vérifié.

## 3. Créer un événement

Dans la section événements, choisissez créer, donnez nom/type/date et fuseau, puis ajoutez une ou plusieurs cérémonies avec lieu et horaire. Vérifiez les dates avant publication.

## 4. Ajouter les invités

Ajoutez un invité à la fois ou utilisez l’import disponible dans l’espace invité. Téléchargez le modèle Excel proposé par l’application, remplissez colonnes identifiées, passez par l’aperçu/mapping, corrigez erreurs puis validez le commit. Ne mettez que les renseignements nécessaires.

## 5. Tables et placement

Configurez le plan par cérémonie, tables/zones/capacité et placez les invités autorisés. Vérifiez invités sans place et capacité avant continuer.

## 6. Choisir et préparer le design

Sélectionnez ou créez un design depuis l’événement. Les fonctions d’assistance IA dépendent du fournisseur et du quota actif; relisez toute proposition avant de l’appliquer.

## 7. Prévisualiser et générer

Relisez noms, cérémonies, lieu, dates et RSVP; démarrez le lot uniquement après confirmation. La génération réserve des crédits; le résultat dépend du nombre d’invités et des états de rendu. Téléchargez PDF individuel ou ZIP du lot disponible.

## 8. Crédits et achat

Le Wallet affiche des **crédits d’usage, pas de l’argent**. Consultez le pack et son prix avant checkout. L’acceptation CGV/politique remboursement applicable doit être affichée par le flux. Le vrai fournisseur externe demeure à vérifier; ne suivez que la méthode indiquée par le site officiel.

## 9. Réponse et accès

L’invité ouvre son lien signé pour voir la partie publiée et répondre cérémonie par cérémonie. À l’accueil, l’organisateur autorisé scanne le QR si appareil/navigateur le permet, sinon saisit le token. Une deuxième entrée pour même invité/cérémonie est limitée par l’application.

## 10. Profil, pages légales et aide

Gérez profil et préférences depuis compte; consultez CGU, confidentialité, vente/remboursement et pages disponibles dans l’app. Contact opérateur: email/téléphone ci-dessus.
""")

DOCUMENTS["10-user-guides/ADMIN_GUIDE.md"] = doc("Guide administrateur InvitaFlow", """
`apps/admin` expose la navigation vers les vues Web qui existent réellement. Les README source indiquent vues `/admin` support/modération, `/admin/pricing` finance pricing et `/admin/finance` paiements; audit/moderation requièrent Support/Super Admin ou Finance Admin selon route. Les contrôles exacts viennent controllers et UI, pas seulement le menu.

## Accès et tâches

1. Authentifiez un compte Keycloak doté du rôle spécifique; n’accordez pas Super Admin pour tester.
2. Support: ouvrir les vues d’audit/modération existantes, examiner l’état et enregistrer une décision autorisée.
3. Finance: consulter pricing schedules/paiements/anomalies selon rôle.
4. Les ajustements wallet passent par operations internes idempotentes et auditables, jamais SQL direct.
5. Avant agir sur payment provider, confirmer provider/ref/status server side; jamais copier secrets marchand dans notes.

Pas de gestion admin utilisateur/stockage/système complète prouvée par le shell admin seul. Captures non générées; voir `screenshots/README.md`.
""")

DOCUMENTS["10-user-guides/AGENCY_GUIDE.md"] = doc("Guide agence InvitaFlow", """
Le schéma Events contient `AgencyWorkspace`, membership/roles, clients/client events, subscriptions et quota reservation. Ces fonctions sont **PARTIAL** : il faut vérifier chaque endpoint/UI et droits par compte avant mise en service.

Guide de base: agence crée workspace si route/permission disponible; owner/admin ajoute membre avec rôle disponible; associe client/événement; vérifie quota/abonnement avant opération; suit rapport réellement exposé. Ne pas supposer équipe multi-workspace ou renouvellement automatique s’il n’est pas exposé et testé.

Captures réelles non générées. Propriétaire doit compléter ce guide après walkthrough vérifié de l’espace Agency.
""", "PARTIAL")

DOCUMENTS["10-user-guides/PARTNER_GUIDE.md"] = doc("Guide partenaire InvitaFlow", """
Payments comporte des modèles Partner, Attribution, Commission Ledger, Payout et routes/services partenaires. Attribution/commissions sont partielles; payout `PAID` ne doit pas être considéré comme virement réel sans fournisseur de versement et confirmation documentés.

Ne communiquer aucun taux, délai, méthode payout ou garantie de revenu sans conditions partenaires approuvées et mise en œuvre vérifiée. Vérifier code d’attribution et tableau réellement visible, puis consulter support pour litige. Captures non générées; fonctionnalités PARTIAL.
""", "PARTIAL")

DOCUMENTS["11-legal/README.md"] = doc("Index documentation juridique", """
Les textes légaux versionnés sont dans `apps/web/content/legal/`, notamment CGU, politique confidentialité, CGV, politique remboursement/annulation, mentions légales et conditions agence/partenaire. Un package manifeste/version se trouve dans `packages/legal-contract/`.

Ces documents du dépôt sont des versions produit et doivent être approuvés par conseil juridique pour le territoire/activité. L’acceptation enregistre version et timestamp; elle ne signifie pas que la conformité est auditée.

Liens runtime au site: `/legal/cgu`, `/legal/confidentialite` et pages associées qui existent sous `apps/web/src/app/`.
""")

DOCUMENTS["00-product/SRS.md"] = doc("Spécification des exigences logiciel (SRS)", """
## Contrôle de version

Identifiant: IV-SRS-1.0 · Date: 2026-10-07 · Source: dépôt courant. Les exigences marquées **observées** décrivent le code et non une validation d'acceptation ou de production. Toute exigence à confirmer demeure en attente de validation produit.

## Objectif et frontières

InvitaFlow aide des organisateurs à préparer des événements, gérer des invités, produire/distribuer des invitations et suivre les réponses. Le système inclut une web app Next.js, des APIs derrière un gateway, des services de domaine, et des dépendances d'identité, données, messagerie et fichiers. Les paiements externes, SMS, payout et intégrations dont les contrats marchands sont absents ne sont pas considérés disponibles.

## Acteurs

| ID | Acteur | Besoin |
|---|---|---|
| ACT-01 | Organisateur connecté | gérer événement, cérémonies, invités, design, invitations et RSVP |
| ACT-02 | Invité public | ouvrir une invitation et accepter/refuser selon le lien disponible |
| ACT-03 | Membre agence | opérations client/workspace selon rôles effectivement accordés |
| ACT-04 | Partenaire | consulter attribution/commission lorsqu’exposées |
| ACT-05 | Support/modération | traiter les dossiers autorisés et tracer les décisions |
| ACT-06 | Finance | suivre prix/paiements selon permission |
| ACT-07 | Worker/service | exécuter travaux asynchrones via identités de service |
| ACT-08 | Identity provider | authentifier les utilisateurs et appliquer les actions email |

## Exigences fonctionnelles

| ID | Exigence | Priorité | Preuve / état |
|---|---|---|---|
| FR-AUTH-001 | L’utilisateur s’authentifie par OIDC Authorization Code avec PKCE et callback serveur. | Must | Web/Keycloak code observé; navigateur runtime à valider |
| FR-AUTH-002 | La session Web n’est acceptée qu’avec email et `email_verified=true`. | Must | callback observé |
| FR-LEGAL-001 | L’inscription accepte les versions légales présentées et enregistre acceptation côté serveur. | Must | action Keycloak et package legal observés; parcours réel à valider |
| FR-EVENT-001 | Un organisateur peut créer et gérer événements/cérémonies selon ownership. | Must | contrôleurs/modèles observés |
| FR-GUEST-001 | Les invités peuvent être gérés/importés et assignés à des événements. | Must | services Guests/Seating observés |
| FR-DESIGN-001 | Le système peut préparer des designs et une assistance éditoriale/design. | Should | services présents; capacité IA distante variable |
| FR-MEDIA-001 | Les médias téléversés sont vérifiés/scannés avant disponibilité. | Must | pipeline code observé; dépendances runtime à valider |
| FR-INV-001 | La génération d’invitations est asynchrone et persistée; son état reste consultable. | Must | service/outbox/worker observés |
| FR-RSVP-001 | RSVP est enregistré par invitation/cérémonie; QR/check-in est vérifié côté serveur. | Must | code/modèles observés; usage réel à valider |
| FR-PAY-001 | Achat de crédits exige confirmation serveur d’un paiement vérifié. | Must | protections Payments/Wallet observées |
| FR-WALLET-001 | Une livraison d’événement dupliquée ne crédite pas le Wallet deux fois. | Must | consumer idempotent et contrainte observés |
| FR-NOTIFY-001 | Des notifications sont persistées et dispatchées selon canaux disponibles. | Should | service Notifications; délivrabilité à valider |
| FR-SEC-001 | Les endpoints appliquent authentification, rôle et/ou ownership pertinents. | Must | guards observés, audit exhaustif à réaliser |
| FR-ADMIN-001 | Support et finance disposent de vues autorisées pour opérations support/pricing. | Should | routes/écrans observés; couverture UI à confirmer |
| FR-AGENCY-001 | Agence/partenaire exposent les fonctions présentes sans présumer des payouts automatisés. | Could | périmètres explicitement partiels |

## Exigences non fonctionnelles

| ID | Exigence vérifiable | Priorité | Validation requise |
|---|---|---|---|
| NFR-PERF-001 | Les pages SSR rendent dates/heures sans dépendre du fuseau implicite de l’hôte. | Must | audit et tests ciblés |
| NFR-OBS-001 | Les erreurs et traces portent une corrélation de requête sans données/secrets sensibles. | Must | tests d’intégration et revue logs |
| NFR-SEC-001 | Les transitions financières sont atomiques, append-only et idempotentes. | Must | tests transactionnels/concurrence |
| NFR-SEC-002 | Les fichiers restent privés jusqu’à passage des contrôles du pipeline. | Must | vérifier ACL et URLs signées runtime |
| NFR-SEC-003 | Les données sont isolées par service et migration possédée par son service. | Must | vérifier déploiement réel |
| NFR-ACCESS-001 | Les thèmes clair et sombre préservent contraste lisible selon WCAG AA. | Should | audit contrastes automatisé + humain |
| NFR-ACCESS-002 | Les layouts sont utilisables sur viewport mobile et desktop. | Should | appareils réels / navigateur |
| NFR-AVAIL-001 | Les travaux asynchrones sont rejouables/diagnostiquables et ne perdent pas les demandes après transaction. | Must | tests de panne et broker |
| NFR-SEC-004 | Les secrets de production ne sont ni commités ni écrits dans les logs. | Must | scan et revue de configuration |
| NFR-AVAIL-002 | Les objectifs RTO/RPO et disponibilité sont formellement approuvés avant tout SLA. | Open | décision produit/opérations; aucune cible approuvée attestée |
| NFR-PRIV-001 | La collecte reste limitée aux données nécessaires; rétention/suppression doivent être approuvées par policy. | Must | revue code + politique/conseil juridique requise |
| NFR-PERF-002 | Un objectif p95 et charge maximale est défini avant engagement de performance. | Open | aucun benchmark/SLO approuvé attesté |

## Contraintes et hypothèses

PostgreSQL est partitionné par service/schema; RabbitMQ transporte les événements durables; Redis et MinIO ont des usages séparés. Keycloak est IdP. Kinshasa est fuseau souhaité pour le runtime RDC mais chaque périmètre de date doit être confirmé. Aucun seuil SLA ou volume officiel n’est attesté par le dépôt.

## Traçabilité

Les FR/NFR renvoient aux exigences/données/API/parcours dans `FUNCTIONAL_REQUIREMENTS.md`, `NON_FUNCTIONAL_REQUIREMENTS.md`, `../01-architecture/`, `../03-api/`, `../05-payments/`, `../06-auth-email/` et `../07-testing/`. La colonne preuve différencie le code observé des essais d’acceptation.
""")

DOCUMENTS["02-data/DATA_DICTIONARY.md"] = doc("Dictionnaire de données généré depuis Prisma", """
Ce fichier est généré depuis les déclarations `model`, `enum` et `field` présentes dans les schémas Prisma source au moment du build documentaire. Il ne prétend pas décrire les données effectivement déployées. Les relations Prisma sont des liens déclaratifs; les migrations doivent être vérifiées sur la base cible.

Le générateur inscrit ci-dessous une entrée par schéma, modèle, champ, enum et valeur. Voir la source des champs liée à chaque schéma pour le sens métier; la présence d’un champ ne certifie ni chiffrement au repos, ni rétention, ni complétude.
""")

DOCUMENTS["09-release/ARCHITECTURE_AUDIT.md"] = doc("Audit architecture — constats et limites", """
## Périmètre

Lecture statique des sources et README présents dans le dépôt au 2026-10-07. Aucun Docker, environnement staging/prod, navigateur authentifié ni base déployée n’a été interrogé pendant la production documentaire.

## Observations

- Monorepo avec frontends séparés, gateway API et services de domaine NestJS/Prisma.
- Les frontières de données et ownership par service sont documentés dans `02-data/DATABASE_OWNERSHIP.md`.
- Les échanges asynchrones utilisent outbox, RabbitMQ et consommateurs idempotents dans plusieurs domaines.
- Paiement réel non démontré: le README Payments déclare FlexPay bloqué sans contrat marchand; ancien nom EasyPay ne constitue pas preuve d’intégration opérationnelle.
- L’application admin est actuellement un shell orientant vers l’interface Web; pas de preuve d’un second système admin indépendant.
- Les modèles agency/partner existent, avec fonctionnalités partielles.
- Les procédures sont documentées mais aucun exercice de restauration ou incident n’est attesté ici.

## Risques et suites

Confirmer contrats externes; exécuter authentification/registration/email en staging; vérifier migrations et configuration réelle; tester pannes RabbitMQ/MinIO/Mailpit/SMTP; mesurer reprise sauvegarde; obtenir approbation juridique; réaliser audit complet rôles/ownership/tenant; fixer SLO/capacité; capturer parcours navigateur réel.
""", "STATIC AUDIT ONLY")

DOCUMENTS["09-release/SECURITY_GAPS.md"] = doc("Écarts et contrôles sécurité — revue statique", """
## Non évalué / ouvert

- Aucun audit de pénétration, SAST/DAST complet ou test de charge n’a été réalisé par cette passe.
- La posture effective TLS, pare-feu, segmentation, IAM cloud, KMS, ACL MinIO, rotation secret et rétention dépend du déploiement et n’est pas déduite du code.
- La couverture exhaustive des autorisations et accès cross-tenant exige une revue route par route et tests adversariaux.
- Vérifier CORS, CSP, cookies/session, rate limits et journaux depuis environnement réel.
- FlexPay/EasyPay et remboursements doivent suivre les contrats et preuves de callback vérifiés; pas de bypass frontend.
- Les documents juridiques doivent être validés par conseil compétent.

## Contrôles présents observés

OIDC/PKCE, vérification email côté callback, protection CSRF/state/nonce, guards bearer/roles, accès par propriétaire dans des handlers, ledger append-only, contraintes uniques/idempotence, upload quarantiné/scan, HMAC QR et audit sont visibles dans le code. Leur présence ne certifie pas l’absence de vulnérabilité.

Prioriser: tests d’IDOR/tenant, vérification signature IPN et replay, fuzz parser fichiers, scan dépendances/containers, session/cookie review, pentest indépendant, DR rehearsal, tests d’alertes et revue de secrets.
""", "OPEN ITEMS")

DOCUMENTS["09-release/OPEN_ITEMS.md"] = doc("Registre de travail ouvert", """
1. **P0 — Exploitation**: confirmer configuration SMTP, DNS/DMARC/TLS, sauvegarde chiffrée hors hôte, restore exercé et alertes on-call.
2. **P0 — Paiements**: obtenir contrat officiel fournisseur RDC et clés sandbox; vérifier signature, retries, reconciliation, refund et rapprochement avant activation.
3. **P0 — Sécurité**: revue systématique ownership/tenant; session config et stockage; scan dépendances/images; test intrusion externe.
4. **P1 — Identité**: parcours réel nouvelle inscription, vérification email, email_verified ID token, reset, expiration et révocation de session.
5. **P1 — Documents**: générer captures après environnement authentifié et vérifier mobile/desktop sur navigateurs réels.
6. **P1 — Agence/partenaire**: définir permissions, conditions commerciales, payout/reconciliation et UX avant annoncer les fonctions.
7. **P1 — Fiabilité**: tester queue DLQ, outbox retries, fichiers volumineux/corrompus, worker crash/reprise.
8. **P2 — Qualité**: audit accessibilité et contraste dans les pages finales, notamment tableaux/badges/états vides.
9. **P2 — SLO**: fixer seuils perf/disponibilité, capacité, RTO/RPO, rétention et responsables avant engagement.

Ce registre décrit des actions recommandées à partir d’une lecture statique; les priorités doivent être réévaluées par les responsables.
""", "OPEN ITEMS")

DOCUMENTS["08-operations/PRODUCTION_READINESS.md"] = doc("Liste de contrôle préparation production", """
## Identité et Web

- [ ] Client/redirect URI de production exacts et scopes `openid profile email`; PKCE S256.
- [ ] Claims email/email_verified inclus; test inscription → vérification → callback réel.
- [ ] Cookies Secure/SameSite, HTTPS, CSP/CORS, retour d’URL interne uniquement.

## Mail et paiements

- [ ] Local Docker par défaut `MAIL_PROVIDER=mailpit`; SMTP réel explicite et secrets hors Git.
- [ ] SPF/DKIM/DMARC et délivrabilité de domaine vérifiés.
- [ ] Fournisseur RDC officiel et contrat marchand approuvés; sandbox/IPN/signature/refund/reconciliation testés avant paiement live.

## Données et reprise

- [ ] Migrations inspectées et appliquées par service avec plan de retour.
- [ ] Sauvegarde chiffrée, séparée du host, restore exercé.
- [ ] ACL privés MinIO, scan antivirus, rétention et suppression vérifiés.
- [ ] RabbitMQ durable, DLQ surveillée, outbox backlog visible; Redis limite/rôle documentés.

## Opérations

- [ ] TLS/secret manager, comptes minimaux, rotation et redaction logs.
- [ ] Health/readiness, alerting, on-call, corrélation et playbooks testés.
- [ ] Capacité/performance/accessibilité testées avec objectifs et appareil/navigateur consignés.
- [ ] Incident response, support, conformité juridique approuvés.

**État de cette passe: NOT VERIFIED.** Les cases ne sont pas cochées par simple présence documentaire.
""", "NOT VERIFIED")

ADR = {
"0001-microservices": ("Architecture de services métier", "Séparer les capacités métier en services propriétaires avec routes/API dédiées.", "Permet déploiement et ownership par domaine, au prix de coordination, événements et observabilité distribuée.", "`services/*`, Gateway et monorepo.", "ACCEPTED — architecture observée"),
"0002-keycloak": ("Keycloak pour OIDC", "Utiliser Keycloak comme Identity Provider et le Web comme BFF OIDC/PKCE.", "Concentre identité et actions email; exige une configuration realm/client cohérente en runtime.", "realm export, client Web, callback BFF.", "ACCEPTED — code statique"),
"0003-rabbitmq": ("RabbitMQ pour travaux et événements durables", "Utiliser RabbitMQ pour flux asynchrones où les domaines le configurent.", "Découple producteurs/consommateurs; retry et DLQ doivent être supervisés.", "Compose, definitions et publishers/consumers.", "ACCEPTED — par domaine"),
"0004-database-per-service": ("Base par service", "Chaque domaine possède son schéma Prisma et son magasin relationnel provisionné.", "Évite jointures interservices; demande projection ou événement pour lectures cross-domain.", "`services/*/prisma`, PostgreSQL provisioning.", "ACCEPTED — observé"),
"0005-outbox-inbox": ("Outbox transactionnelle et déduplication consommateur", "Persister le message lié à l’état métier dans la transaction; dédupliquer l’effet côté consumer si mécanisme présent.", "Protège le dual-write et retries; patterns et couverture varient selon service.", "Outbox modules, tables et consumers.", "PARTIAL — vérifier domaine par domaine"),
"0006-payment-provider-abstraction": ("Abstraction provider paiement", "Isoler les fournisseurs derrière PaymentProvider et utiliser Mock en local.", "Permet simuler checkout; une intégration réelle attend contrat et test sandbox.", "Payments README et provider adapters.", "ACCEPTED — FlexPay réel bloqué"),
"0007-easypay": ("EasyPay adapter non qualifié", "Conserver la référence adapter/code historique sans le qualifier de fournisseur prêt pour production.", "Préserve l’implémentation à investiguer tout en empêchant une promesse ou un live non validé.", "README Payments, provider Easypay et `docs/PAYMENTS_EASYPAY.md`.", "NOT VERIFIED — contrat sandbox/merchant requis"),
"0008-wallet-credits": ("Wallet en unités de crédit", "Traiter le solde comme des unités d’usage inscrites dans un ledger append-only; ce n’est pas de l’argent.", "Permet déduplication et audit; correction financière par opération compensatoire.", "Wallet schema/service et payment consumers.", "ACCEPTED — code observé"),
"0009-minio-storage": ("MinIO pour objets privés", "Stocker médias et artefacts privés dans MinIO avec séparation quarantaine/ready quand applicable.", "Protège l’accès par contrôle et URL limitée; ACL et restore doivent être vérifiés en runtime.", "Media/Invitations et infrastructure MinIO.", "ACCEPTED — code, runtime à valider"),
"0010-design-document-v2": ("Document de design versionné", "Représenter le design invitation comme document JSON versionné, distinct du rendu final.", "Sépare édition et artefact; compatibilité renderer et migration de schema à préserver.", "`packages/design-document` et services Designs/Rendering.", "ACCEPTED — code observé"),
"0011-legal-registration": ("Acceptation légale au parcours d’inscription", "Valider l’acceptation et enregistrer côté serveur les versions/timestamps du package présenté.", "Crée une trace explicite; le contenu et sa validité juridique nécessitent un conseil compétent.", "Keycloak legal form action, legal-contract package.", "ACCEPTED — parcours runtime à valider"),
"0012-email-verification": ("Vérification email obligatoire", "Keycloak exige VERIFY_EMAIL et la session Web refuse un claim d’email non vérifié.", "Réduit comptes non qualifiés; dépend du bon scope ID token et délivrabilité réelle.", "Realm export, provisioning email, callback.", "ACCEPTED — tests E2E à faire"),
}

DIAGRAMS = [
  ("system-context","flow","Contexte système",["Organisateur","Invité","FOCUS HD / InvitaFlow","Keycloak","Email provider","Opérateur support"]),
  ("c4-containers","flow","C4 conteneurs",["Browser","Next.js Web BFF","Gateway","Domain services","PostgreSQL par domaine","RabbitMQ / Redis / MinIO"]),
  ("component-architecture","flow","Composants",["Web routes","BFF auth/session","Gateway proxy","Domain API","Outbox & workers","Infrastructure"]),
  ("deployment","flow","Déploiement logique",["Client browser","Traefik edge","Web + Gateway","Services API/workers","Data plane","External IdP/mail/provider"]),
  ("network-trust-boundary","flow","Frontières de confiance",["Internet untrusted","Reverse proxy TLS","Authenticated web/API","Private services","DB/broker/object store","External providers"]),
  ("authentication-sequence","sequence","Authentification OIDC",["Browser","Next.js BFF","Keycloak","Session Redis","Gateway"]),
  ("registration-legal-acceptance","sequence","Inscription et acceptation légale",["Visitor","Keycloak form","Legal package","Email action","Mailbox","Web callback"]),
  ("email-verification-sequence","sequence","Vérification email",["New user","Keycloak","Mail provider","Verification URL","Web callback"]),
  ("password-reset-sequence","sequence","Réinitialisation mot de passe",["User","Keycloak","Email provider","Reset page","New session"]),
  ("event-creation-sequence","sequence","Création événement",["Organizer","Web BFF","Gateway","Events API","Events DB","Outbox/Broker"]),
  ("guest-import-sequence","sequence","Import invités",["Organizer","Web","Guests API","Import parser/job","PostgreSQL","Notifications"]),
  ("invitation-generation-sequence","sequence","Génération invitation",["Organizer","Invitations API","Wallet","Outbox/RabbitMQ","Rendering worker","MinIO"]),
  ("design-document-rendering-pipeline","flow","Pipeline rendu design",["Design version","Renderer","Template + assets","PDF output","Validation","Private artifact"]),
  ("pdf-zip-generation-sequence","sequence","PDF/ZIP batch",["Web client","Invitations API","RabbitMQ","Worker","MinIO","Download endpoint"]),
  ("rsvp-sequence","sequence","RSVP public",["Guest","Public link","Invitations API","RSVP DB","Notification outbox"]),
  ("qr-generation-sequence","flow","QR signé",["Invitation + ceremony","Canonical payload","HMAC signature","QR image","Printed invitation"]),
  ("check-in-sequence","sequence","Check-in",["Organizer device","Web scanner","Events/Invitation API","Signature validator","Unique check-in DB"]),
  ("anti-double-entry","flow","Prévention double check-in",["Scan arrives","Verify event/ceremony","Atomic unique insert","Already checked?","Record one attendance","Return conflict/duplicate"]),
  ("payment-checkout-sequence","sequence","Checkout paiement",["User","Web/Gateway","Payments API","Provider adapter","Provider hosted checkout","Database/outbox"]),
  ("easypay-ipn-sequence","sequence","Callback fournisseur (à valider)",["Provider IPN","Payments callback","Signature/reference check","Payment transaction","Outbox","Wallet consumer"]),
  ("payment-idempotency-sequence","sequence","Idempotence paiement",["Provider retries","Callback handler","Unique receipt/ref","Compare-and-set status","Anomaly/replay result"]),
  ("wallet-credit-sequence","sequence","Crédit Wallet",["Payment verified","payment.succeeded event","Wallet consumer","Unique event ledger","Balance projection"]),
  ("agency-flow","flow","Flux agence",["Agency admin","Workspace","Members + roles","Client / events","Quota check","Reports where available"]),
  ("partner-flow","flow","Flux partenaire",["Partner attribution","Referral/event","Commission ledger","Payout request/state","Settlement provider (not verified)","Reconciliation"]),
  ("media-upload-storage-sequence","sequence","Upload média sécurisé",["Browser","Presigned upload","Quarantine bucket","Media API/ClamAV","Decode/transcode","Private ready bucket"]),
  ("ai-design-composer-sequence","sequence","AI Design Composer",["Organizer","Web","AI Design API","Quota/advisory lock","RabbitMQ worker","Model/preview storage"]),
  ("erd-high-level","erd","ERD haut niveau",["Profile","Event/Ceremony","Guest/Group","Design/Version","Invitation/Batch/RSVP","Payment/Wallet/Ledger","Media/Asset","Audit/Notification"]),
  ("payment-state-machine","state","Machine paiement",["CREATED","PENDING","PROCESSING","SUCCEEDED","FAILED","CANCELLED","EXPIRED","REFUND_PENDING","REFUNDED"]),
  ("invitation-state-machine","state","Machine invitation batch",["RESERVING","QUEUED","GENERATING","COMPLETED","PARTIAL","FAILED","CANCELLED"]),
  ("event-domain-lifecycle","state","Cycle événement",["DRAFT","PUBLISHED","CANCELLED","COMPLETED"]),
  ("backup-restore-architecture","flow","Sauvegarde et restauration",["Databases","Consistent backup job","Encrypted offsite storage","Restore isolation","Migration/validation","Service cutover"]),
  ("observability-architecture","flow","Observabilité",["Web/Gateway/services","Structured logs","Metrics/health","Correlation IDs","Dashboards/alerts","Operator runbook"]),
]

SEQUENCE_MESSAGES = {
"Authentification OIDC": ["GET /login", "Create PKCE state", "Authenticate at Keycloak", "Exchange code + verifier", "Validate claims / session"],
"Inscription et acceptation légale": ["Open registration", "Submit legal versions", "Persist consent timestamp", "Create unverified user", "Send verify-email link"],
"Vérification email": ["Require verification", "Send one-time link", "Validate action link", "Set email verified", "Issue verified claim"],
"Réinitialisation mot de passe": ["Request reset", "Issue expiring link", "Send reset email", "Set new password", "Authenticate again"],
"Création événement": ["Submit event", "Forward bearer request", "Validate owner / fields", "Commit event", "Publish outbox event"],
"Import invités": ["Upload spreadsheet", "Parse rows / size", "Persist import job", "Deduplicate and upsert", "Report row errors"],
"Génération invitation": ["Create batch", "Reserve credits once", "Commit batch / outbox", "Render PDF items", "Store private outputs"],
"PDF/ZIP batch": ["Request owned export", "Check status / expiry", "Fetch private objects", "Stream download", "Record access result"],
"RSVP public": ["Open signed link", "Validate signature", "Submit ceremony RSVP", "Persist response once", "Notify organizer"],
"Check-in": ["Scan QR", "Resolve event / ceremony", "Verify signature / owner", "Insert unique attendance", "Return result / duplicate"],
"Checkout paiement": ["Request checkout", "Validate order state", "Create provider session", "Persist provider refs", "Return checkout URL"],
"Callback fournisseur (à valider)": ["Receive provider IPN", "Resolve provider ref", "Verify sig / amount", "Commit state + outbox", "Credit wallet once"],
"Idempotence paiement": ["Receive callback", "Validate provider payload", "Insert unique receipt", "Compare-and-set status", "Return prior / anomaly"],
"Crédit Wallet": ["Verify success event", "Publish versioned event", "Deduplicate event ID", "Append credit ledger", "Update balance view"],
"Upload média sécurisé": ["Request scoped upload", "Upload to quarantine", "Verify bytes / size", "Scan and transcode", "Promote private object"],
"AI Design Composer": ["Submit AI job", "Check owner / quota", "Persist and enqueue", "Generate proposal", "Store result / status"],
}

DOCX_EXPORTS = {
"01_Vision_Produit_InvitaFlow.docx": "00-product/PRODUCT_VISION.md",
"02_Cahier_Des_Charges_SRS_InvitaFlow.docx": "00-product/SRS.md",
"03_Architecture_InvitaFlow.docx": "01-architecture/ARCHITECTURE.md",
"04_Architecture_Securite_InvitaFlow.docx": "04-security/SECURITY_ARCHITECTURE.md",
"05_Modele_Donnees_InvitaFlow.docx": "02-data/ERD.md",
"06_API_Integration_InvitaFlow.docx": "03-api/API_GUIDE.md",
"07_Architecture_Paiement_InvitaFlow.docx": "05-payments/PAYMENT_PROVIDER_ARCHITECTURE.md",
"08_Strategie_Tests_InvitaFlow.docx": "07-testing/TEST_STRATEGY.md",
"09_Runbook_Exploitation_InvitaFlow.docx": "08-operations/RUNBOOK.md",
"10_Deployment_InvitaFlow.docx": "08-operations/DEPLOYMENT.md",
"11_Backup_Disaster_Recovery_InvitaFlow.docx": "08-operations/BACKUP_RECOVERY.md",
"12_Release_Readiness_InvitaFlow.docx": "09-release/RELEASE_READINESS_REPORT.md",
"13_Guide_Utilisateur_InvitaFlow.docx": "10-user-guides/END_USER_GUIDE.md",
"14_Guide_Administrateur_InvitaFlow.docx": "10-user-guides/ADMIN_GUIDE.md",
"15_Guide_Agence_InvitaFlow.docx": "10-user-guides/AGENCY_GUIDE.md",
"16_Guide_Partenaire_InvitaFlow.docx": "10-user-guides/PARTNER_GUIDE.md",
}
DOCX_DIAGRAMS = {
"03_Architecture_InvitaFlow.docx": ["system-context", "c4-containers"],
"04_Architecture_Securite_InvitaFlow.docx": ["network-trust-boundary", "authentication-sequence"],
"05_Modele_Donnees_InvitaFlow.docx": ["erd-high-level"],
"07_Architecture_Paiement_InvitaFlow.docx": ["payment-state-machine", "wallet-credit-sequence"],
"10_Deployment_InvitaFlow.docx": ["deployment"],
"11_Backup_Disaster_Recovery_InvitaFlow.docx": ["backup-restore-architecture"],
}

def ensure_dirs() -> None:
    for rel in DOCUMENTS:
        (DOCS / rel).parent.mkdir(parents=True, exist_ok=True)

def write_source_docs() -> None:
    for rel, content in DOCUMENTS.items():
        path = DOCS / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        if not path.exists():
            path.write_text(content, encoding="utf-8")
    for slug, (title, decision, consequence, evidence, status) in ADR.items():
        body = f"""# ADR {slug[:4]} — {title}

{META}
Status: **{status}**

## Contexte

Le produit doit conserver un choix d’architecture explicite et traçable à partir du dépôt courant.

## Décision

{decision}

## Conséquences

{consequence}

## Alternatives considérées

Le dépôt ne conserve pas de trace formelle des alternatives contemporaines. Cette ADR est rédigée rétrospectivement depuis le code; l’alternative à réévaluer lors d’une nouvelle revue est un choix plus centralisé ou plus simple à opérer. Ce paragraphe ne prétend pas établir qu’une telle analyse a été réalisée historiquement.

## Sources

{evidence}
"""
        path = DOCS / "adr" / f"{slug}.md"
        path.parent.mkdir(parents=True, exist_ok=True)
        if not path.exists():
            path.write_text(body, encoding="utf-8")

    write_data_dictionary()
    write_diagrams()
    write_screenshot_material()
    write_indexes()
    write_adr_index_and_reports()

def write_data_dictionary() -> None:
    rows = ["# Dictionnaire de données Prisma", "", META, "", "Schémas source présents dans le dépôt. Les noms et champs sont extraits des fichiers Prisma versionnés; ils ne démontrent pas l’état d’une base déployée.", ""]
    schemas = sorted(ROOT.glob("services/*/prisma/schema.prisma"))
    for schema in schemas:
        domain = schema.parts[-3]
        text = schema.read_text(encoding="utf-8", errors="replace")
        rows += [f"## {domain}", "", f"Source: `{schema.relative_to(ROOT).as_posix()}`", ""]
        for kind, name, body in re.findall(r"\b(model|enum)\s+(\w+)\s*\{([^}]*)\}", text, flags=re.S):
            rows += [f"### {kind.title()} `{name}`", ""]
            if kind == "model":
                rows += ["| Champ | Type Prisma | Obligatoire | Unique | Description métier | Service propriétaire | Sensible ? | Rétention |", "|---|---|---|---|---|---|---|---|"]
                for line in body.splitlines():
                    stripped = line.strip()
                    if not stripped or stripped.startswith("//") or stripped.startswith("@@"):
                        continue
                    match = re.match(r"(\w+)\s+([^\s]+)(.*)", stripped)
                    if match and not match.group(1).startswith("@@"):
                        field, typ, attrs = match.groups()
                        # Nullable and uniqueness are inferred only from Prisma syntax.
                        attrs = attrs.split("//", 1)[0].strip()
                        required = "Oui" if not typ.endswith("?") else "Non"
                        unique = "Oui" if "@id" in attrs or "@unique" in attrs else "Non"
                        sensitive = "Oui — revue requise" if re.search(r"email|phone|mobile|token|secret|password|address|notes|subject|reference|ip", field, re.I) else "À classer"
                        rows.append(f"| `{field}` | `{typ}` {html.escape(attrs)} | {required} | {unique} | Non documentée dans Prisma | `{domain}` | {sensitive} | Politique approuvée non trouvée |")
                rows.append("")
            else:
                values = [line.strip() for line in body.splitlines() if re.match(r"\s*\w+\s*$", line) and not line.strip().startswith("//")]
                rows += [", ".join(f"`{x}`" for x in values), ""]
    if not schemas:
        rows += ["Aucun schéma Prisma trouvé — génération à corriger.", ""]
    (DOCS / "02-data/DATA_DICTIONARY.md").write_text("\n".join(rows), encoding="utf-8")

def mermaid_source(kind: str, title: str, nodes: list[str]) -> str:
    if kind == "sequence":
        ids = [f"P{i+1}" for i in range(len(nodes))]
        parts = ["sequenceDiagram", "    autonumber"]
        parts.extend(f"    participant {ident} as {label}" for ident, label in zip(ids, nodes))
        messages=SEQUENCE_MESSAGES.get(title, [f"{title} step {i+1}" for i in range(len(ids)-1)])
        for i, message in enumerate(messages[:len(ids)-1]):
            parts.append(f"    {ids[i]}->>{ids[i+1]}: {message}")
        if len(ids) > 2:
            parts.append(f"    {ids[-1]}-->>{ids[0]}: Résultat / statut vérifié")
        return "\n".join(parts) + "\n"
    if kind == "state":
        parts = ["stateDiagram-v2", f"    title: {title}"]
        parts.extend(f"    {a} --> {b}" for a, b in state_edges(title))
        return "\n".join(parts) + "\n"
    if kind == "erd":
        parts = ["erDiagram", "    %% Schéma conceptuel : domaines séparés, aucune relation SQL inter-base représentée."]
        for i, name in enumerate(nodes):
            entity = re.sub(r"\W+", "_", name).strip("_").upper() or f"DOMAIN_{i+1}"
            parts.append(f"    {entity} {{\n        string id\n        string ownerSubject\n    }}")
        return "\n".join(parts) + "\n"
    parts = ["flowchart LR"]
    parts.extend(f"    N{i}[\"{node}\"]" for i, node in enumerate(nodes))
    parts.extend(f"    N{i} --> N{i+1}" for i in range(len(nodes) - 1))
    return "\n".join(parts) + "\n"

def state_edges(title: str) -> list[tuple[str,str]]:
    if title.startswith("Machine paiement"):
        return [("[*]","CREATED"),("CREATED","PROCESSING"),("CREATED","CANCELLED"),("PROCESSING","PENDING"),("PROCESSING","CREATED"),("PENDING","PROCESSING"),("PENDING","SUCCEEDED"),("PENDING","FAILED"),("PENDING","CANCELLED"),("PENDING","EXPIRED"),("PROCESSING","SUCCEEDED"),("PROCESSING","FAILED"),("PROCESSING","CANCELLED"),("PROCESSING","EXPIRED"),("SUCCEEDED","REFUND_PENDING"),("REFUND_PENDING","REFUNDED"),("REFUND_PENDING","SUCCEEDED"),("REFUNDED","[*]"),("FAILED","[*]"),("CANCELLED","[*]"),("EXPIRED","[*]")]
    if title.startswith("Machine invitation"):
        return [("[*]","RESERVING"),("RESERVING","QUEUED"),("RESERVING","FAILED"),("RESERVING","CANCELLED"),("QUEUED","GENERATING"),("QUEUED","CANCELLED"),("GENERATING","COMPLETED"),("GENERATING","PARTIAL"),("GENERATING","FAILED"),("GENERATING","CANCELLED"),("PARTIAL","QUEUED"),("PARTIAL","[*]"),("COMPLETED","[*]"),("FAILED","[*]"),("CANCELLED","[*]")]
    return [("[*]","DRAFT"),("DRAFT","PUBLISHED"),("DRAFT","CANCELLED"),("DRAFT","COMPLETED"),("PUBLISHED","CANCELLED"),("PUBLISHED","COMPLETED"),("CANCELLED","[*]"),("COMPLETED","[*]")]

def draw_diagram(title: str, nodes: list[str], kind: str, svg_path: Path, png_path: Path) -> None:
    # Same simple node/edge representation is exported to SVG and PNG so no CLI renderer is required.
    width = 1400
    wrapped = [textwrap.wrap(label, 24) or [label] for label in nodes]
    box_h = max(96, max(sum(25 for _ in lines) + 34 for lines in wrapped))
    if kind == "sequence":
        width = max(1400, len(nodes) * 205)
        height = 660
        n = len(nodes)
        gap = width / (n + 1)
        centers = [int(gap * (i + 1)) for i in range(n)]
        boxes = [(x - 82, 70, 164, 66, label) for x, label in zip(centers, nodes)]
        svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">', '<rect width="100%" height="100%" fill="#fbf7f0"/>', f'<text x="40" y="42" font-size="27" font-family="Arial" font-weight="700" fill="#301338">{html.escape(title)}</text>']
        for x, y, w, h, label in boxes:
            svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="13" fill="#fff" stroke="#59305d" stroke-width="3"/>')
            svg.append(f'<text x="{x+w/2}" y="{y+39}" text-anchor="middle" font-size="17" font-family="Arial" fill="#29152f">{html.escape(label)}</text>')
            svg.append(f'<line x1="{x+w/2}" y1="{y+h}" x2="{x+w/2}" y2="{height-52}" stroke="#9d8da0" stroke-dasharray="8 7"/>')
        ys = [220 + (i % 4) * 85 for i in range(len(nodes)-1)]
        messages=SEQUENCE_MESSAGES.get(title, [f"{title} {i+1}" for i in range(len(nodes)-1)])
        for i, y in enumerate(ys):
            x1, x2 = centers[i], centers[i+1]
            svg.append(f'<line x1="{x1}" y1="{y}" x2="{x2}" y2="{y}" stroke="#b58924" stroke-width="3" marker-end="url(#arrow)"/>')
            message=messages[i] if i<len(messages) else f"{title} {i+1}"
            svg.append(f'<text x="{(x1+x2)/2}" y="{y-8}" text-anchor="middle" font-size="14" font-family="Arial" fill="#46324a">{html.escape(message[:68])}</text>')
        svg.insert(1, '<defs><marker id="arrow" markerWidth="10" markerHeight="8" refX="9" refY="4" orient="auto"><path d="M0,0 L10,4 L0,8 z" fill="#b58924"/></marker></defs>')
        svg.append('</svg>')
        im = Image.new("RGB", (width, height), "#fbf7f0")
        d = ImageDraw.Draw(im)
        font = ImageFont.load_default()
        d.text((38, 24), title, fill="#301338", font=font)
        for x, y, w, h, label in boxes:
            d.rounded_rectangle((x, y, x+w, y+h), radius=13, fill="white", outline="#59305d", width=3)
            d.text((x+8, y+24), label[:24], fill="#29152f", font=font)
            d.line((x+w//2, y+h, x+w//2, height-50), fill="#9d8da0", width=2)
        for i, y in enumerate(ys):
            x1, x2 = centers[i], centers[i+1]
            d.line((x1, y, x2-7, y), fill="#b58924", width=3)
            d.polygon([(x2-7,y-6),(x2,y),(x2-7,y+6)], fill="#b58924")
            message=messages[i] if i<len(messages) else title
            d.text(((x1+x2)//2-65, y-20), message[:43], fill="#46324a", font=font)
    else:
        cols = (4 if len(nodes) > 8 else 3)
        rows = (len(nodes) + cols - 1) // cols
        width = max(width, cols * 340 + 80)
        height = max(250, 170 + rows * 145)
        positions = [(100 + (i % cols)*340, 105 + (i//cols)*145) for i in range(len(nodes))]
        edge_pairs = [(i, i+1) for i in range(len(nodes)-1)]
        if kind == "state":
            # State exports display meaningful branches instead of implying every state is a successor.
            index={name:i for i,name in enumerate(nodes)}
            positions=[(110+(i%4)*330,110+(i//4)*135) for i in range(len(nodes))]
            edge_pairs=[(index[a],index[b]) for a,b in state_edges(title) if a in index and b in index]
            width=max(1400, 4*330+180); height=max(330,170+((len(nodes)+3)//4)*135)
        svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">', '<defs><marker id="arrow" markerWidth="10" markerHeight="8" refX="9" refY="4" orient="auto"><path d="M0,0 L10,4 L0,8 z" fill="#b58924"/></marker></defs>', '<rect width="100%" height="100%" fill="#fbf7f0"/>', f'<text x="40" y="45" font-size="28" font-family="Arial" font-weight="700" fill="#301338">{html.escape(title)}</text>']
        for i,j in edge_pairs:
            x,y=positions[i]; nx,ny=positions[j]
            # Edge goes in layout order; vertical on row changes.
            if y == ny and nx>x:
                x1, y1, x2, y2 = x+250, y+40, nx-10, ny+40
            elif nx == x and ny>y:
                x1, y1, x2, y2 = x+125, y+95, nx+125, ny-8
            else:
                x1,y1,x2,y2=x+125,y+46,nx+125,ny+46
            svg.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="#b58924" stroke-width="3" marker-end="url(#arrow)"/>')
        for i, ((x,y), lines) in enumerate(zip(positions, wrapped)):
            fill = "#f1e8d8" if i % 2 == 0 else "#fff"
            if kind == "state" and i in (0, len(nodes)-1):
                fill = "#ead9a9"
            svg.append(f'<rect x="{x}" y="{y}" width="250" height="92" rx="15" fill="{fill}" stroke="#59305d" stroke-width="3"/>')
            for j, line in enumerate(lines[:3]):
                svg.append(f'<text x="{x+125}" y="{y+36+j*22}" text-anchor="middle" font-size="16" font-family="Arial" fill="#29152f">{html.escape(line)}</text>')
        svg.append('</svg>')
        im = Image.new("RGB", (width, height), "#fbf7f0")
        d = ImageDraw.Draw(im)
        f = ImageFont.load_default()
        d.text((40,25), title, fill="#301338", font=f)
        # Draw edges first so they stay behind the nodes.
        for i,j in edge_pairs:
            x,y=positions[i]; nx,ny=positions[j]
            if y==ny and nx>x: x1,y1,x2,y2=x+250,y+46,nx,ny+46
            elif nx==x and ny>y: x1,y1,x2,y2=x+125,y+92,nx+125,ny
            else: x1,y1,x2,y2=x+125,y+46,nx+125,ny+46
            d.line((x1,y1,x2,y2),fill="#b58924",width=3)
        for i, ((x,y), lines) in enumerate(zip(positions, wrapped)):
            fill = "#f1e8d8" if i%2==0 else "white"
            if kind == "state" and i in (0, len(nodes)-1): fill = "#ead9a9"
            d.rounded_rectangle((x,y,x+250,y+92), radius=15, fill=fill, outline="#59305d", width=3)
            for j,line in enumerate(lines[:3]): d.text((x+12,y+25+j*20),line[:34],fill="#29152f",font=f)
    png_path.parent.mkdir(parents=True, exist_ok=True)
    svg_path.parent.mkdir(parents=True, exist_ok=True)
    svg_path.write_text("\n".join(svg), encoding="utf-8")
    im.save(png_path, "PNG", optimize=True)

def write_diagrams() -> None:
    index = ["# Architecture diagrams", "", META, "", "Chaque diagramme est indépendant : Mermaid éditable dans `sources/`, export SVG et PNG. SVG/PNG sont generated with a small deterministic renderer from the node/edge model; they preserve the communication path but do not execute Mermaid CLI. Mermaid sources are the editable source of truth.", ""]
    for slug, kind, title, nodes in DIAGRAMS:
        source = DOCS / "01-architecture/diagrams/sources" / f"{slug}.mmd"
        source.parent.mkdir(parents=True, exist_ok=True)
        source.write_text(mermaid_source(kind,title,nodes),encoding="utf-8")
        svg = DOCS / "01-architecture/diagrams/svg" / f"{slug}.svg"
        png = DOCS / "01-architecture/diagrams/png" / f"{slug}.png"
        draw_diagram(title,nodes,kind,svg,png)
        index += [f"## {title}", "", f"![{title}](diagrams/png/{slug}.png)", "", f"Source: [`{slug}.mmd`](diagrams/sources/{slug}.mmd) · [`SVG`](diagrams/svg/{slug}.svg) · [`PNG`](diagrams/png/{slug}.png)", ""]
    (DOCS / "01-architecture/DIAGRAM_CATALOG.md").write_text("\n".join(index),encoding="utf-8")

def write_screenshot_material() -> None:
    entries = [
      ("01-public-home","Accueil public","/","Tout visiteur","01-public-home.png","Noter viewport desktop/mobile et thème"),
      ("02-registration","Inscription","Keycloak / route runtime à relever","Visiteur fictif","02-registration.png","Masquer adresse du compte de test"),
      ("03-legal-consent","Consentement légal","Keycloak registration form","Visiteur fictif","03-legal-consent.png","Afficher versions des textes et cases"),
      ("04-email-verification","Vérification email","Lien d’action Keycloak","Compte fictif","04-email-verification.png","Boîte de test uniquement"),
      ("05-login","Connexion","Keycloak login / URL runtime à relever","Compte fictif","05-login.png","Ne jamais montrer tokens ou URL à usage unique"),
      ("06-dashboard","Dashboard","/dashboard","Organisateur","06-dashboard.png","Données de démonstration clairement fictives"),
      ("07-create-event","Création événement","/events","Organisateur","07-create-event.png","Formulaire réel"),
      ("08-event-ceremonies","Cérémonies","/events/:eventId","Organisateur propriétaire","08-event-ceremonies.png","Utiliser événement fictif"),
      ("09-guests","Invités","/events/:eventId/guests","Organisateur propriétaire","09-guests.png","Noms/adresses fictifs"),
      ("10-guest-import-guide","Guide import invités","/events/:eventId/guests","Organisateur propriétaire","10-guest-import-guide.png","Montrer modèle si disponible"),
      ("11-excel-import-preview","Aperçu import Excel","/events/:eventId/guests","Organisateur propriétaire","11-excel-import-preview.png","Fichier de fixture synthétique"),
      ("12-seating","Placement","/events/:eventId/seating","Organisateur propriétaire","12-seating.png","Fixture de sièges"),
      ("13-design-editor","Éditeur design","/events/:eventId/designs","Organisateur propriétaire","13-design-editor.png","Visuel autorisé"),
      ("14-invitation-preview","Aperçu invitation","/events/:eventId/invitations","Organisateur propriétaire","14-invitation-preview.png","Contenu fictif uniquement"),
      ("15-wallet","Wallet","/account/wallet","Organisateur","15-wallet.png","Solde de démonstration"),
      ("16-payment-checkout","Checkout","Provider test uniquement; route à relever","Compte sandbox autorisé","16-payment-checkout.png","Aucun moyen réel ou secret"),
      ("17-rsvp","RSVP","/invite/:token","Invité de démonstration","17-rsvp.png","Token temporaire de fixture, flouter après parcours"),
      ("18-checkin","Check-in","/events/:eventId/check-in","Organisateur/agent autorisé","18-checkin.png","QR de démonstration"),
      ("19-contact","Contact","/contact","Tout visiteur","19-contact.png","Aucune conversation client réelle"),
      ("20-legal-index","Index légal","/legal","Tout visiteur","20-legal-index.png","Textes publiés et version à noter"),
    ]
    checklist = ["# Checklist captures UI utilisateur", "", META, "", "État de cette livraison : **NOT GENERATED**. La stack UI n’a pas été démarrée ni validée dans un navigateur authentifié durant la passe de documentation. Chaque capture ci-dessous attend une preuve runtime réelle; aucun mockup ne la remplace.", "", "| ID | Scène attendue | Statut | Fichier attendu | Notes de capture |", "|---|---|---|---|---|"]
    for i, (identifier,name,route,role,filename,notes) in enumerate(entries,1):
        checklist.append(f"| {identifier} | End-user | {name} | `{route}` | Desktop + Mobile | Light + Dark | {role} | NOT GENERATED | `{filename}` | {notes} |")
    checklist += ["", "Ne pas capturer PII client, tokens, paramètres secrets, écran d’administration contenant des données réelles ni QR de production. Documenter commit/app build, browser+OS, dimensions, thème et parcours. Export PNG originaux."]
    end_user = DOCS / "10-user-guides/screenshots/end-user"
    end_user.mkdir(parents=True, exist_ok=True)
    (end_user / "SCREENSHOT_CHECKLIST.md").write_text("# Checklist des captures du guide utilisateur\n\nLa checklist canonique des 20 captures utilisateur et 6 captures admin est [`../../SCREENSHOT_CHECKLIST.md`](../../SCREENSHOT_CHECKLIST.md). Toutes les lignes sont actuellement **NOT GENERATED**. Cette page sert uniquement de redirection pour les références historiques.\n",encoding="utf-8")
    root_checklist=["# Plan de captures documentaires", "", META, "", "Aucune capture n’a été générée. Chaque capture doit provenir d’un navigateur/runtime réel avec données de démonstration fictives et sans secrets/PII. Les routes marquées runtime à relever doivent être confirmées avant capture.", "", "## End-user", "", "| ID | Guide | Screen | Route | Desktop/Mobile | Light/Dark | Required account/role | Status | Filename | Notes |", "|---|---|---|---|---|---|---|---|---|---|"]
    for identifier,name,route,role,filename,notes in entries:
        root_checklist.append(f"| {identifier} | End-user | {name} | `{route}` | Desktop + Mobile | Light + Dark | {role} | NOT GENERATED | `{filename}` | {notes} |")
    root_checklist += ["", "## Admin — pages existantes", "", "| ID | Guide | Screen | Route | Desktop/Mobile | Light/Dark | Required account/role | Status | Filename | Notes |", "|---|---|---|---|---|---|---|---|---|---|"]
    admins=[("ADM-01","Console support/modération","/admin","Support ou Super Admin","admin-01-moderation.png"),("ADM-02","Finance / paiements","/admin/finance","Finance Admin","admin-02-finance.png"),("ADM-03","Tarification","/admin/pricing","Finance Admin","admin-03-pricing.png"),("ADM-04","Partenaires","/admin/partners","Finance Admin","admin-04-partners.png"),("ADM-05","Stockage","/admin/storage","rôle autorisé par route","admin-05-storage.png"),("ADM-06","Analytics","/admin/analytics","rôle administrateur","admin-06-analytics.png")]
    for ident,name,route,role,filename in admins:
        root_checklist.append(f"| {ident} | Admin | {name} | `{route}` | Desktop + Mobile | Light + Dark | {role} | NOT GENERATED | `{filename}` | Capture seulement après vérification du rôle effectif. |")
    (DOCS/"10-user-guides/SCREENSHOT_CHECKLIST.md").write_text("\n".join(root_checklist)+"\n",encoding="utf-8")
    for role in ("admin","agency","partner"):
        dest = DOCS / "10-user-guides/screenshots" / role
        dest.mkdir(parents=True, exist_ok=True)
        (dest/"README.md").write_text(f"# Captures {role}\n\nStatut : **NOT GENERATED**. Parcours à capturer uniquement après validation runtime réelle avec compte fictif autorisé. Voir le guide [`../{role.upper()}_GUIDE.md`](../../{role.upper()}_GUIDE.md) et `../end-user/SCREENSHOT_CHECKLIST.md`. Ne jamais placer de PII client ou secret.\n",encoding="utf-8")
    for role in ("end-user","admin","agency","partner"):
        visual=DOCS/"10-user-guides/visual"/role; visual.mkdir(parents=True,exist_ok=True)
        (visual/"README.md").write_text(f"# Visuels annotés — {role}\n\nAucun visuel généré : les captures applicatives réelles sont **NOT GENERATED** faute de session navigateur/runtime authentifiée vérifiée. Ne pas fabriquer de visuels UI. Dès capture autorisée, déposer l’original et une annotation qui ne contient aucune donnée réelle.\n",encoding="utf-8")

def write_adr_index_and_reports() -> None:
    adr_dir=DOCS/"adr"; adr_dir.mkdir(parents=True,exist_ok=True)
    adr_lines=["# Architecture Decision Records", "", "Décisions fondées sur le dépôt. Les statuts restent descriptifs; ils ne remplacent pas une approbation formelle d’architecture.", ""]
    for slug,(title,_decision,_consequence,_evidence,status) in ADR.items():
        adr_lines.append(f"- [ADR {slug[:4]} — {title}]({slug}.md) — {status}")
    (adr_dir/"README.md").write_text("\n".join(adr_lines)+"\n",encoding="utf-8")
    reports=DOCS/"reports"; reports.mkdir(parents=True,exist_ok=True)
    report_docs={
      "DOCUMENTATION_AUDIT.md": doc("Audit de la documentation", "Corpus statique produit depuis le dépôt du 2026-10-07. Le validateur contrôle liens locaux, fichiers requis, 32 jeux de sources Mermaid et exports PNG/SVG, 16 DOCX canoniques (intégrité ZIP, XML, titre, marque et contenu), les 26 lignes de captures NOT GENERATED et un scan de secrets y compris les DOCX historiques archivés. Les exports PNG/SVG sont des rendus déterministes du générateur de projet; Mermaid CLI validation reste NOT VERIFIED. Les Word n’ont pas reçu de QA visuelle : LibreOffice est absent. Les captures runtime sont NOT GENERATED; runtime evidence reste PENDING. Architecture, sécurité, disponibilité, performance et E2E ne sont pas certifiés. Aucune stack applicative n’a été lancée pendant la génération documentaire.", "STATIC DOCUMENTATION AUDIT"),
      "ARCHITECTURE_AUDIT.md": (DOCS/"09-release/ARCHITECTURE_AUDIT.md").read_text(encoding="utf-8"),
      "SECURITY_GAPS.md": (DOCS/"09-release/SECURITY_GAPS.md").read_text(encoding="utf-8"),
      "OPEN_ITEMS.md": (DOCS/"09-release/OPEN_ITEMS.md").read_text(encoding="utf-8"),
    }
    for filename,content in report_docs.items(): (reports/filename).write_text(content,encoding="utf-8")

def write_indexes() -> None:
    groups = {
      "00-product":"Vision, périmètre, SRS, exigences, règles métier",
      "01-architecture":"Architecture, C4, 32 diagrammes, ADR",
      "02-data":"ERD, ownership, dictionnaire Prisma, classification et cycle de vie",
      "03-api":"Guide/catalogue API, erreurs, événements, intégration",
      "04-security":"Sécurité, STRIDE, autorisations, secrets, incidents et tests",
      "05-payments":"Providers, paiement, EasyPay/FlexPay état réel, Wallet et preuves",
      "06-auth-email":"Keycloak, inscription, acceptation légale, email et SMTP",
      "07-testing":"Stratégie, E2E, sécurité, performance, mobile, accessibilité",
      "08-operations":"Runbook, deploy, migrations, sauvegarde, observabilité, readiness",
      "09-release":"ADR/audits, release, risques, rollback, open items",
      "10-user-guides":"guides utilisateurs/admin/agence/partenaire et captures attendues",
      "11-legal":"pointeurs vers le package juridique versionné",
    }
    lines=["# Documentation technique InvitaFlow", "", META, "", "Corpus de référence construit depuis le dépôt local. Les documents marquent explicitement l’implémentation observée, partielle, planifiée ou non vérifiée. Une description de code ne certifie pas une disponibilité runtime.", "", "## Navigation"]
    for folder,description in groups.items():
        lines += [f"### [{folder}](./{folder}/)", "", description+".", ""]
    lines += ["## Artefacts", "", "- [32 diagrammes source Mermaid et exports](01-architecture/DIAGRAM_CATALOG.md)", "- [ADRs](adr/README.md)", "- [Rapports d’audit](reports/)", "- [Export DOCX (16 documents)](dist/docx/README.md)", "- [Plan central des captures](10-user-guides/SCREENSHOT_CHECKLIST.md)", "- [Guide utilisateur](10-user-guides/END_USER_GUIDE.md)", "- [Archive historique](_archive/README.md)", "- [Outillage de génération et validation](../tools/docs/README.md)", "", "## Documentation status", "", "- SOURCE DOCUMENTATION: READY", "- DIAGRAM SOURCES: READY", "- PNG/SVG EXPORTS: GENERATED", "- MERMAID CLI VALIDATION: NOT VERIFIED", "- WORD EXPORTS: GENERATED", "- WORD STRUCTURAL VALIDATION: PASS", "- WORD VISUAL QA: NOT VERIFIED", "- END-USER SCREENSHOTS: NOT GENERATED", "- ADMIN SCREENSHOTS: NOT GENERATED", "- RUNTIME EVIDENCE: PENDING", "", "Les diagrammes PNG/SVG sont des exports déterministes du générateur documentaire. Les fichiers Mermaid sont éditables; l’export ne remplace pas la validation par Mermaid CLI.", "", "## Documents existants conservés", "", "Les dossiers `docs/` antérieurs couvrent déjà parcours AI, design, expérience client, intégration historique EasyPay, validations E2E et opérations/sauvegarde. Ils demeurent référentiels complémentaires : [E2E historique](E2E_VALIDATION.md), [remaining work](REMAINING_WORK.md), [backup operations](OPERATIONS_BACKUPS.md), [AI customer journey](AI_CUSTOMER_JOURNEY.md), [AI design composer](AI_DESIGN_COMPOSER.md), [AI editorial assist](AI_EDITORIAL_ASSIST.md), [historique EasyPay](PAYMENTS_EASYPAY.md).", "", "## Support", "", "Exploitant propriétaire: FOCUS HD ENTREPRISES · `fucushd098@gmail.com` · `+243973431495`. Ne pas envoyer secrets ou données d’invités sur ce canal.", ""]
    (DOCS/"README.md").write_text("\n".join(lines),encoding="utf-8")
    dist=DOCS/"dist/guides"; dist.mkdir(parents=True,exist_ok=True)
    (dist/"README.md").write_text("# Distribution des guides\n\nLes guides Markdown éditables et les captures réelles demeurent dans `docs/10-user-guides/`. Aucune capture n’étant générée dans cette passe, ce dossier ne contient pas de copie visuelle fictive. Les versions Word se trouvent dans `docs/dist/docx/`.\n",encoding="utf-8")
    diagrams=DOCS/"dist/diagrams"; diagrams.mkdir(parents=True,exist_ok=True)
    (diagrams/"README.md").write_text("# Distribution des diagrammes\n\nLes 32 sources Mermaid et exports séparés sont conservés dans `docs/01-architecture/diagrams/`. Ce dossier d’index évite de dupliquer les exports.\n",encoding="utf-8")

def add_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading=OxmlElement("w:shd"); shading.set(qn("w:fill"),fill); tc_pr.append(shading)

def add_markdown_to_doc(document, content: str) -> None:
    lines=content.splitlines(); i=0
    while i<len(lines):
        line=lines[i]
        if not line.strip(): i+=1; continue
        if line.startswith("|") and i+1<len(lines) and re.match(r"\s*\|?[-: |]+\|\s*$",lines[i+1]):
            matrix=[]
            while i<len(lines) and lines[i].startswith("|"):
                cells=[x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+",x or " ") for x in cells): matrix.append(cells)
                i+=1
            if matrix:
                t=document.add_table(rows=0,cols=max(map(len,matrix))); t.style="Light Shading Accent 1"
                for row_index,row in enumerate(matrix):
                    cells=t.add_row().cells
                    for j,val in enumerate(row):
                        cells[j].text=re.sub(r"`([^`]+)`",r"\1",val).replace("**","")
                        cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        if row_index==0: add_cell_shading(cells[j],"EEE7F0")
            continue
        if line.startswith("# "):
            i+=1; continue
        if line.startswith("## "):
            document.add_heading(line[3:],level=1); i+=1; continue
        if line.startswith("### "):
            document.add_heading(line[4:],level=2); i+=1; continue
        if line.startswith("#### "):
            document.add_heading(line[5:],level=3); i+=1; continue
        if line.startswith("```xml") or line.startswith("```powershell") or line.startswith("```text") or line.startswith("```mermaid") or line.startswith("```"):
            i+=1; block=[]
            while i<len(lines) and not lines[i].startswith("```"): block.append(lines[i]); i+=1
            p=document.add_paragraph(); r=p.add_run("\n".join(block)); r.font.name="Consolas"; r.font.size=Pt(8)
            i+=1; continue
        if line.startswith("- [ ]") or line.startswith("- [x]") or line.startswith("- "):
            p=document.add_paragraph(style="List Bullet"); p.add_run(re.sub(r"`([^`]+)`",r"\1",line.lstrip("- ")).replace("**","")); i+=1; continue
        if re.match(r"\d+\. ",line):
            p=document.add_paragraph(style="List Number"); p.add_run(re.sub(r"`([^`]+)`",r"\1",re.sub(r"^\d+\. ","",line)).replace("**","")); i+=1; continue
        para=[]
        while i<len(lines) and lines[i].strip() and not lines[i].startswith(("#","- ","|","```")) and not re.match(r"\d+\. ",lines[i]):
            para.append(lines[i]); i+=1
        raw=" ".join(para)
        raw=re.sub(r"\[([^]]+)\]\(([^)]+)\)",r"\1",raw); raw=re.sub(r"`([^`]+)`",r"\1",raw); raw=raw.replace("**","")
        document.add_paragraph(raw)

def write_docx() -> None:
    out=DOCS/"dist/docx"; out.mkdir(parents=True,exist_ok=True)
    readme=["# Export Word éditable", "", "16 documents `.docx` générés du Markdown de référence avec python-docx. Pour reconstruire depuis la racine, utiliser l’environnement du dossier tools/docs et lancer `python tools/docs/build_documentation.py`. Les mises à jour se font d’abord dans les sources Markdown puis régénération.", ""]
    for output,source in DOCX_EXPORTS.items():
        content=(DOCS/source).read_text(encoding="utf-8")
        d=Document()
        sec=d.sections[0]; sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
        normal=d.styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(11); normal.font.color.rgb=RGBColor(45,35,47)
        for sty,size,col in (("Title",26,(0,0,0)),("Heading 1",18,(72,35,78)),("Heading 2",14,(94,52,96)),("Heading 3",11,(140,101,43))):
            s=d.styles[sty]; s.font.name="Aptos Display"; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(*col)
        title=re.search(r"^# (.+)$",content,re.M)
        p=d.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; p.add_run(title.group(1) if title else Path(source).stem)
        p=d.add_paragraph(); p.add_run("INVITAFLOW  ·  FOCUS HD ENTREPRISES").bold=True; p.runs[0].font.color.rgb=RGBColor(155,119,44)
        p=d.add_paragraph(); p.add_run(f"Document de référence · {TODAY} · Version 1.0").italic=True
        p=d.add_paragraph(); p.add_run("Contact : fucushd098@gmail.com · +243 973 431 495").font.size=Pt(9)
        d.add_paragraph()
        d.add_page_break()
        # Keep main body compact and readable; the title/meta repeats are filtered.
        filtered=re.sub(r"^# .+\n", "", content, count=1)
        filtered=re.sub(r"^Version:.*?(?=^## |\Z)", "", filtered, count=1, flags=re.S|re.M)
        add_markdown_to_doc(d,filtered)
        if Path(output).name in DOCX_DIAGRAMS:
            d.add_heading("Diagrammes de référence", level=1)
            for slug in DOCX_DIAGRAMS[Path(output).name]:
                png=DOCS/"01-architecture/diagrams/png"/f"{slug}.png"
                if png.is_file():
                    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(png), width=Inches(6.15))
        for section in d.sections:
            footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
            run=footer.add_run("InvitaFlow · FOCUS HD  |  Documentation technique  |  "); run.font.size=Pt(8); run.font.color.rgb=RGBColor(105,91,107)
            fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); footer._p.append(fld)
        path=out/output; path.parent.mkdir(parents=True,exist_ok=True); d.save(path)
        readme.append(f"- `{Path(output).name}` — `{source}`")
    (out/"README.md").write_text("\n".join(readme)+"\n",encoding="utf-8")

def main() -> None:
    ensure_dirs()
    write_source_docs()
    write_docx()
    print(f"Wrote {len(DOCUMENTS)} Markdown sources, {len(ADR)} ADRs, {len(DIAGRAMS)} diagrams and {len(DOCX_EXPORTS)} DOCX exports under {DOCS}")

if __name__ == "__main__":
    main()
